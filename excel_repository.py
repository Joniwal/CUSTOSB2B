from __future__ import annotations

import io
import hashlib
import logging
import os
import re
import threading
import unicodedata
import uuid
import zipfile
from calendar import monthrange
from copy import copy
from collections import Counter
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable


HEADERS = [
    "ID",
    "Tipo de Atividade",
    "Status",
    "Situação",
    "Material Utilizado",
    "Cod. Material",
    "Quantidade",
    "Custo Mat",
    "Serviço M.O",
    "Qtde Serviço",
    "Custo M.O",
    "Custo Total",
    "Custo Evitado",
    "Custo Técnico / Dia",
    "Qtde Técnicos",
    "Qtde Dias",
    "Custo por Técnico",
    "GAP",
    "Atualizado Em",
    "Data",
    "Empresa",
    "Nome do Técnico",
    "Matrícula",
    "EPS",
    "Região",
    "Cluster",
    "Tecnologia",
    "Ganho Esperado/Importância",
    "DRAFT",
]

FIELD_NAMES = [
    "id",
    "tipo_atividade",
    "status",
    "situacao",
    "material_utilizado",
    "cod_material",
    "quantidade",
    "custo_mat",
    "servico_mo",
    "qtde_servico",
    "custo_mo",
    "custo_total",
    "custo_evitado",
    "custo_tecnico_dia",
    "qtde_tecnicos",
    "qtde_dias",
    "custo_por_tecnico",
    "gap",
    "atualizado_em",
    "data",
    "empresa",
    "tecnico_nome",
    "matricula",
    "eps",
    "regiao",
    "cluster",
    "tecnologia",
    "ganho_esperado",
    "draft",
]

NUMERIC_FIELDS = {
    "quantidade",
    "custo_mat",
    "qtde_servico",
    "custo_mo",
    "custo_total",
    "custo_evitado",
    "custo_tecnico_dia",
    "qtde_tecnicos",
    "qtde_dias",
    "custo_por_tecnico",
    "gap",
}

# Nomes equivalentes encontrados nas bases operacionais já existentes.
# O primeiro nome de cada campo continua sendo o cabeçalho padrão do projeto.
EXTRA_HEADER_ALIASES = {
    "tipo_atividade": ("Obra Executada", "Tipo Atividade"),
    "status": ("Status Obra",),
    "situacao": ("Detalhes Obra", "Ganho Esperado/Importancia"),
    "custo_mat": ("Custo Material",),
    "qtde_tecnicos": ("Vol. Técnicos", "Volume Técnicos"),
    "qtde_dias": ("Tempo Execução",),
    "data": ("Data", "Atualizado Em"),
    "tecnico_nome": ("Nome do tecnico", "Técnico", "Tecnico"),
    "matricula": ("Matricula",),
    "tecnologia": ("BBR/B2B/GPON",),
    "ganho_esperado": ("Ganho Esperado/Importancia",),
    "draft": ("Draft",),
}

SERVICE_CATEGORIES = (
    {"key": "implantacao", "label": "Implantação", "tokens": ("implant",)},
    {"key": "reparo", "label": "Reparo", "tokens": ("reparo",)},
    {"key": "ativacao", "label": "Ativação", "tokens": ("ativacao",)},
)

DEFAULT_TECHNOLOGIES = (
    "ERB",
    "GPON",
    "VISTORIA",
    "REPARO",
    "ATIVAÇÃO",
    "MAGIC TOOLS",
    "RETIRADA EQUIPAMENTOS",
)

SERVICE_CALC_SHEET = "Calculo_Servicos"
SERVICE_CALC_HEADERS = [
    "Chave Registro",
    "ID Atividade",
    "Chave Serviço",
    "Código Serviço",
    "Serviço",
    "Unidade",
    "Valor Unitário",
    "Quantidade",
    "Subtotal",
    "Atualizado Em",
]
SERVICE_HEADER_ALIASES = {
    "code": ("CODIGO", "CÓDIGO", "COD", "COD. SERVICO", "COD. SERVIÇO"),
    "description": ("SERVICO", "SERVIÇO", "DESCRICAO", "DESCRIÇÃO"),
    "unit_price": (
        "CUSTO_UNITARIO",
        "CUSTO UNITARIO",
        "CUSTO UNITÁRIO",
        "VALOR UNITARIO",
        "VALOR UNITÁRIO",
        "PRECO",
        "PREÇO",
        "VALOR",
    ),
    "unit": ("UNIDADE", "UN", "UND"),
}

MATERIAL_CALC_SHEET = "Calculo_Materiais"
MATERIAL_CALC_HEADERS = [
    "Chave Registro",
    "ID Atividade",
    "Chave Material",
    "Código Material",
    "Material",
    "Unidade",
    "Valor Unitário",
    "Quantidade",
    "Subtotal",
    "Atualizado Em",
]
MATERIAL_HEADER_ALIASES = {
    "code": (
        "MAT_CODE",
        "MAT CODE",
        "CODIGO",
        "CÓDIGO",
        "COD",
        "COD. MATERIAL",
        "COD MATERIAL",
    ),
    "description": ("MATERIAL", "DESCRICAO", "DESCRIÇÃO", "ITEM"),
    "unit_price": (
        "VALOR",
        "CUSTO_UNITARIO",
        "CUSTO UNITARIO",
        "CUSTO UNITÁRIO",
        "VALOR UNITARIO",
        "VALOR UNITÁRIO",
        "PRECO",
        "PREÇO",
    ),
    "unit": ("UNIDADE", "UN", "UND"),
}


class RepositoryError(RuntimeError):
    pass


log = logging.getLogger(__name__)


def normalize_header(value: Any) -> str:
    text = unicodedata.normalize("NFKD", str(value or ""))
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"[^a-z0-9]+", "", text.casefold())


def resolve_field_columns(headers: Iterable[Any]) -> dict[str, int]:
    """Map canonical fields to zero-based columns, accepting known aliases."""
    available: dict[str, int] = {}
    for index, header in enumerate(headers):
        key = normalize_header(header)
        if key and key not in available:
            available[key] = index

    result: dict[str, int] = {}
    for header, field in zip(HEADERS, FIELD_NAMES):
        aliases = (header, *EXTRA_HEADER_ALIASES.get(field, ()))
        for alias in aliases:
            index = available.get(normalize_header(alias))
            if index is not None:
                result[field] = index
                break
    return result


def resolve_service_columns(headers: Iterable[Any]) -> dict[str, int]:
    """Map service catalog headers to zero-based columns."""
    available = {
        normalize_header(header): index
        for index, header in enumerate(headers)
        if normalize_header(header)
    }
    result: dict[str, int] = {}
    for field, aliases in SERVICE_HEADER_ALIASES.items():
        for alias in aliases:
            index = available.get(normalize_header(alias))
            if index is not None:
                result[field] = index
                break
    return result


def resolve_material_columns(headers: Iterable[Any]) -> dict[str, int]:
    """Map material catalog headers to zero-based columns."""
    available = {
        normalize_header(header): index
        for index, header in enumerate(headers)
        if normalize_header(header)
    }
    result: dict[str, int] = {}
    for field, aliases in MATERIAL_HEADER_ALIASES.items():
        for alias in aliases:
            index = available.get(normalize_header(alias))
            if index is not None:
                result[field] = index
                break
    return result


def calculation_mode(field_columns: dict[str, int]) -> str:
    detailed_fields = {
        "quantidade",
        "custo_mat",
        "qtde_servico",
        "custo_mo",
        "custo_tecnico_dia",
        "qtde_tecnicos",
        "qtde_dias",
    }
    return "detailed" if detailed_fields.issubset(field_columns) else "direct_costs"


def to_number(value: Any) -> float:
    if value in (None, ""):
        return 0.0
    if isinstance(value, bool):
        return float(value)
    if isinstance(value, (int, float)):
        return float(value)
    text = str(value).strip().replace("R$", "").replace(" ", "")
    if "," in text and "." in text:
        text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")
    try:
        return float(text)
    except ValueError:
        return 0.0


def clean_text(value: Any, max_length: int = 240) -> str:
    return str(value or "").strip()[:max_length]


def format_date(value: Any) -> str:
    if isinstance(value, datetime):
        return value.isoformat(timespec="minutes")
    if isinstance(value, date):
        return datetime.combine(value, datetime.min.time()).isoformat(timespec="minutes")
    return clean_text(value, 40)


def activity_date(value: Any) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = clean_text(value, 40)
    if not text:
        return None
    for candidate in (text[:10], text):
        for date_format in ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y"):
            try:
                return datetime.strptime(candidate, date_format).date()
            except ValueError:
                continue
    return None


def business_days_for_month(month_text: str) -> int:
    """Count Monday-Friday days; holidays can be adjusted in Configurações."""
    try:
        selected = datetime.strptime(clean_text(month_text, 7), "%Y-%m")
    except ValueError:
        selected = datetime.combine(date.today().replace(day=1), datetime.min.time())
    days = monthrange(selected.year, selected.month)[1]
    return sum(1 for day in range(1, days + 1) if date(selected.year, selected.month, day).weekday() < 5)


def calculate_values(activity: dict[str, Any], *, use_team_gap: bool = True) -> dict[str, float]:
    # The activity form stores consolidated service and material amounts. GAP
    # compares the service amount with the configured cost of the assigned team.
    custo_total = round(to_number(activity.get("custo_mat")) + to_number(activity.get("custo_mo")), 2)
    tecnicos = to_number(activity.get("qtde_tecnicos"))
    custo_por_tecnico = round(custo_total / tecnicos, 2) if tecnicos > 0 else 0.0
    custo_equipe = (
        to_number(activity.get("custo_tecnico_dia"))
        * tecnicos
        * to_number(activity.get("qtde_dias"))
    )
    # GAP positivo: os serviços avaliados superam o custo da equipe.
    # GAP negativo: o custo da equipe supera o valor dos serviços.
    gap = (
        round(to_number(activity.get("custo_mo")) - custo_equipe, 2)
        if use_team_gap
        else round(to_number(activity.get("gap")), 2)
    )
    return {
        "custo_total": custo_total,
        "custo_por_tecnico": custo_por_tecnico,
        "gap": gap,
    }


def normalize_activity(raw: dict[str, Any], *, use_team_gap: bool = True) -> dict[str, Any]:
    result: dict[str, Any] = {}
    for field in FIELD_NAMES:
        value = raw.get(field)
        if field in NUMERIC_FIELDS:
            result[field] = round(to_number(value), 2)
        elif field == "data":
            parsed_date = activity_date(value)
            result[field] = parsed_date.isoformat() if parsed_date else ""
        elif field == "atualizado_em":
            result[field] = format_date(value)
        else:
            result[field] = clean_text(value)
    result["calculation_mode"] = clean_text(raw.get("calculation_mode"), 24) or "detailed"
    calculated = calculate_values(result, use_team_gap=use_team_gap)

    # Numeric totals already present in an operational workbook are authoritative.
    # Formula cells are recalculated from the detailed inputs because data_only=False.
    raw_total = raw.get("custo_total")
    if raw_total not in (None, "") and not str(raw_total).lstrip().startswith("="):
        calculated["custo_total"] = round(to_number(raw_total), 2)
        technicians = to_number(result.get("qtde_tecnicos"))
        raw_per_technician = raw.get("custo_por_tecnico")
        if raw_per_technician not in (None, "") and not str(raw_per_technician).lstrip().startswith("="):
            calculated["custo_por_tecnico"] = round(to_number(raw_per_technician), 2)
        else:
            calculated["custo_por_tecnico"] = (
                round(calculated["custo_total"] / technicians, 2) if technicians > 0 else 0.0
            )
        raw_gap = raw.get("gap")
        if raw_gap not in (None, "") and not str(raw_gap).lstrip().startswith("="):
            calculated["gap"] = round(to_number(raw_gap), 2)
        else:
            calculated["gap"] = calculate_values(result, use_team_gap=use_team_gap)["gap"]

    result.update(calculated)
    return result


def validate_update(
    activity_id: str,
    payload: dict[str, Any],
    *,
    use_team_gap: bool = True,
) -> dict[str, Any]:
    if not isinstance(payload, dict):
        raise ValueError("Dados da atividade inválidos.")
    required = ("tipo_atividade", "status", "situacao")
    missing = [field for field in required if not clean_text(payload.get(field))]
    if missing:
        raise ValueError("Preencha tipo de atividade, status e situação.")

    activity = {"id": clean_text(activity_id, 60)}
    for field in FIELD_NAMES[1:]:
        if field in ("custo_total", "custo_por_tecnico", "gap", "atualizado_em"):
            continue
        activity[field] = to_number(payload.get(field)) if field in NUMERIC_FIELDS else clean_text(payload.get(field))

    for field in NUMERIC_FIELDS:
        if field in activity and activity[field] < 0 and field != "gap":
            raise ValueError("Valores de quantidade e custo não podem ser negativos.")
    activity["atualizado_em"] = datetime.now().isoformat(timespec="minutes")
    activity.update(calculate_values(activity, use_team_gap=use_team_gap))
    return normalize_activity(activity, use_team_gap=use_team_gap)


def aggregate_by(activities: list[dict[str, Any]], field: str) -> list[dict[str, Any]]:
    grouped: dict[str, dict[str, float | int | str]] = {}
    for item in activities:
        label = clean_text(item.get(field)) or "Não informado"
        bucket = grouped.setdefault(
            label,
            {"label": label, "quantity": 0, "labor": 0.0, "material": 0.0, "cost": 0.0, "gap": 0.0},
        )
        bucket["quantity"] = int(bucket["quantity"]) + 1
        bucket["labor"] = round(float(bucket["labor"]) + to_number(item.get("custo_mo")), 2)
        bucket["material"] = round(float(bucket["material"]) + to_number(item.get("custo_mat")), 2)
        bucket["cost"] = round(float(bucket["cost"]) + to_number(item.get("custo_total")), 2)
        bucket["gap"] = round(float(bucket["gap"]) + to_number(item.get("gap")), 2)
    return sorted(grouped.values(), key=lambda item: (-int(item["quantity"]), str(item["label"]).casefold()))


def service_category(value: Any) -> str:
    normalized = normalize_header(value)
    for category in SERVICE_CATEGORIES:
        if any(token in normalized for token in category["tokens"]):
            return str(category["key"])
    return ""


def summarize_service_categories(
    activities: list[dict[str, Any]],
    settings: dict[str, Any],
) -> list[dict[str, Any]]:
    configured_teams = {
        clean_text(item.get("category"), 80): max(0, int(to_number(item.get("technicians"))))
        for item in settings.get("category_teams", [])
        if isinstance(item, dict) and clean_text(item.get("category"), 80)
    }
    monthly_cost = max(0.0, to_number(settings.get("monthly_technician_cost")))
    result: list[dict[str, Any]] = []
    for definition in SERVICE_CATEGORIES:
        key = str(definition["key"])
        label = str(definition["label"])
        category_items = [item for item in activities if service_category(item.get("tipo_atividade")) == key]
        service_total = round(sum(to_number(item.get("custo_mo")) for item in category_items), 2)
        material_total = round(sum(to_number(item.get("custo_mat")) for item in category_items), 2)
        total_cost = round(service_total + material_total, 2)
        technicians = configured_teams.get(key, 0)
        team_value = round(monthly_cost * technicians, 2)
        result.append(
            {
                "key": key,
                "label": label,
                "activity_count": len(category_items),
                "technicians": technicians,
                "team_value": team_value,
                "service": service_total,
                "material": material_total,
                "total": total_cost,
                "gap": round(service_total - team_value, 2),
            }
        )
    return result


def summarize_activities(activities: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(activities)
    concluido_labels = {"concluida", "concluido"}
    cancelado_labels = {"cancelada", "cancelado", "inativa", "inativo"}

    def folded(value: Any) -> str:
        return normalize_header(value)

    concluidas = sum(1 for item in activities if folded(item.get("status")) in concluido_labels)
    canceladas = sum(1 for item in activities if folded(item.get("status")) in cancelado_labels)
    em_andamento = max(0, total - concluidas - canceladas)
    status_counts = Counter(clean_text(item.get("status")) or "Sem status" for item in activities)
    type_counts = Counter(clean_text(item.get("tipo_atividade")) or "Sem categoria" for item in activities)

    recent = sorted(
        activities,
        key=lambda item: item.get("data") or item.get("atualizado_em") or "",
        reverse=True,
    )[:5]
    financial = {
        "custo_mo": round(sum(to_number(item.get("custo_mo")) for item in activities), 2),
        "custo_material": round(sum(to_number(item.get("custo_mat")) for item in activities), 2),
        "custo_total": round(sum(to_number(item.get("custo_total")) for item in activities), 2),
        "custo_evitado": round(sum(to_number(item.get("custo_evitado")) for item in activities), 2),
        "gap": round(sum(to_number(item.get("gap")) for item in activities), 2),
    }
    return {
        "metrics": {
            "total": total,
            "em_andamento": em_andamento,
            "concluidas": concluidas,
            "custo_evitado": financial["custo_evitado"],
        },
        "financial": financial,
        "status_counts": [{"label": key, "value": value} for key, value in status_counts.most_common()],
        "type_counts": [{"label": key, "value": value} for key, value in type_counts.most_common()],
        "charts": {
            "by_type": aggregate_by(activities, "tipo_atividade"),
            "by_technology": aggregate_by(activities, "tecnologia"),
            "by_technician": aggregate_by(activities, "tecnico_nome"),
            "by_company": aggregate_by(activities, "empresa"),
        },
        "recent": recent,
    }


def _deduplicate_existing_directories(candidates: Iterable[Path]) -> list[Path]:
    result: list[Path] = []
    seen: set[str] = set()
    for candidate in candidates:
        try:
            resolved = candidate.expanduser().resolve()
            if not resolved.is_dir():
                continue
        except (OSError, RuntimeError):
            continue
        key = os.path.normcase(str(resolved))
        if key not in seen:
            seen.add(key)
            result.append(resolved)
    return result


def _onedrive_registry_roots() -> list[Path]:
    """Return OneDrive sync roots registered for the current Windows user."""
    if os.name != "nt":
        return []
    try:
        import winreg
    except ImportError:
        return []

    found: list[Path] = []

    def visit(key_path: str, depth: int = 0) -> None:
        try:
            with winreg.OpenKey(winreg.HKEY_CURRENT_USER, key_path) as key:
                for value_name in ("UserFolder", "MountPoint"):
                    try:
                        value, _kind = winreg.QueryValueEx(key, value_name)
                    except OSError:
                        continue
                    if isinstance(value, str) and value.strip():
                        found.append(Path(value.strip()))
                # Team-site libraries can be mounted outside OneDrive*. Their
                # registered cache uses scope IDs as value names, not MountPoint.
                if key_path.rsplit("\\", 1)[-1].casefold() == "scopeidtomountpointpathcache":
                    value_index = 0
                    while True:
                        try:
                            _name, value, _kind = winreg.EnumValue(key, value_index)
                        except OSError:
                            break
                        if isinstance(value, str) and value.strip():
                            candidate = Path(value.strip())
                            if candidate.is_absolute():
                                found.append(candidate)
                        value_index += 1
                if depth >= 3:
                    return
                index = 0
                while True:
                    try:
                        child = winreg.EnumKey(key, index)
                    except OSError:
                        break
                    visit(key_path + "\\" + child, depth + 1)
                    index += 1
        except OSError:
            return

    visit(r"Software\Microsoft\OneDrive\Accounts")
    visit(r"Software\SyncEngines\Providers\OneDrive")
    return found


def _onedrive_search_roots(
    base_dir: Path,
    *,
    include_automatic_when_configured: bool = False,
) -> list[Path]:
    """Use explicit roots or discover personal/corporate OneDrive mounts."""
    configured = os.getenv("EXCEL_SEARCH_ROOTS", "").strip()
    configured_roots: list[Path] = []
    if configured:
        configured_parts = [raw.strip().strip('"') for raw in configured.split(os.pathsep) if raw.strip()]
        placeholder_tokens = ("seu_usuario", "sua empresa", "seu usuário", "<usuario>", "<usuário>")
        has_template_placeholder = any(
            token in raw.casefold()
            for raw in configured_parts
            for token in placeholder_tokens
        )
        if has_template_placeholder:
            log.warning(
                "EXCEL_SEARCH_ROOTS contém valores de exemplo; ignorando a configuração e detectando o OneDrive automaticamente."
            )
            configured_parts = []

        candidates = []
        for raw in configured_parts:
            path = Path(raw)
            candidates.append(path if path.is_absolute() else base_dir / path)
        configured_roots = _deduplicate_existing_directories(candidates)
        if configured_roots and not include_automatic_when_configured:
            return configured_roots
        if configured_parts and not configured_roots:
            log.warning(
                "Nenhuma pasta de EXCEL_SEARCH_ROOTS existe neste computador; detectando o OneDrive automaticamente."
            )

    candidates = [
        Path(value)
        for variable in ("OneDriveCommercial", "OneDriveConsumer", "OneDrive")
        if (value := os.getenv(variable, "").strip())
    ]
    candidates.extend(_onedrive_registry_roots())
    try:
        candidates.extend(path for path in Path.home().glob("OneDrive*") if path.is_dir())
    except OSError:
        pass
    return _deduplicate_existing_directories([*configured_roots, *candidates])


def _sample_fallback_enabled() -> bool:
    return os.getenv("EXCEL_FALLBACK_SAMPLE", "false").strip().casefold() in {"1", "true", "yes", "sim"}


def _read_file_with_shared_access(path: Path) -> bytes:
    """Read a Windows file that Excel/OneDrive currently keeps open."""
    try:
        return path.read_bytes()
    except PermissionError:
        if os.name != "nt":
            raise

    import ctypes
    import msvcrt
    from ctypes import wintypes

    generic_read = 0x80000000
    share_read_write_delete = 0x00000001 | 0x00000002 | 0x00000004
    open_existing = 3
    normal_attributes = 0x00000080

    kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
    create_file = kernel32.CreateFileW
    create_file.argtypes = [
        wintypes.LPCWSTR,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.LPVOID,
        wintypes.DWORD,
        wintypes.DWORD,
        wintypes.HANDLE,
    ]
    create_file.restype = wintypes.HANDLE
    close_handle = kernel32.CloseHandle
    close_handle.argtypes = [wintypes.HANDLE]
    close_handle.restype = wintypes.BOOL

    handle = create_file(
        str(path),
        generic_read,
        share_read_write_delete,
        None,
        open_existing,
        normal_attributes,
        None,
    )
    if handle == ctypes.c_void_p(-1).value:
        raise ctypes.WinError(ctypes.get_last_error())

    try:
        descriptor = msvcrt.open_osfhandle(int(handle), os.O_RDONLY | os.O_BINARY)
    except Exception:
        close_handle(handle)
        raise

    with os.fdopen(descriptor, "rb") as stream:
        return stream.read()


def discover_excel_file(
    filename: str,
    *,
    base_dir: Path,
    fallback_filenames: Iterable[str] = (),
    include_automatic_roots: bool = False,
) -> Path:
    """Find one workbook by prioritized names in synchronized OneDrive folders."""
    filenames: list[str] = []
    seen_names: set[str] = set()
    for raw_name in (filename, *fallback_filenames):
        candidate_name = clean_text(raw_name, 240)
        if not candidate_name:
            continue
        if (
            Path(candidate_name).name != candidate_name
            or candidate_name.startswith("~$")
            or Path(candidate_name).suffix.casefold() not in {".xlsx", ".xlsm"}
        ):
            raise RepositoryError(
                "Os nomes de arquivos Excel devem conter apenas o nome de um arquivo .xlsx ou .xlsm, sem pastas."
            )
        normalized_name = candidate_name.casefold()
        if normalized_name not in seen_names:
            seen_names.add(normalized_name)
            filenames.append(candidate_name)

    if not filenames:
        raise RepositoryError("EXCEL_FILENAME deve informar um arquivo .xlsx ou .xlsm.")

    roots = _onedrive_search_roots(
        base_dir,
        include_automatic_when_configured=include_automatic_roots,
    )
    target_names = {name.casefold() for name in filenames}
    matches: dict[str, dict[str, Path]] = {name.casefold(): {} for name in filenames}
    ignored_directories = {".git", ".pytest_cache", "__pycache__", "node_modules", "backups", "_backups"}

    for root in roots:
        try:
            for current, directories, files in os.walk(root, topdown=True, followlinks=False):
                directories[:] = [
                    directory for directory in directories
                    if directory.casefold() not in ignored_directories
                ]
                for actual_name in files:
                    normalized_actual_name = actual_name.casefold()
                    if actual_name.startswith("~$") or normalized_actual_name not in target_names:
                        continue
                    candidate = Path(current) / actual_name
                    try:
                        resolved = candidate.resolve()
                        if resolved.is_file():
                            matches[normalized_actual_name][os.path.normcase(str(resolved))] = resolved
                    except OSError:
                        continue
        except (OSError, PermissionError):
            log.warning("Não foi possível percorrer a pasta sincronizada: %s", root)

    for candidate_name in filenames:
        ordered = sorted(matches[candidate_name.casefold()].values(), key=lambda path: str(path).casefold())
        if len(ordered) == 1:
            log.info("Excel localizado automaticamente: %s", ordered[0])
            return ordered[0]
        if len(ordered) > 1:
            locations = "\n".join(f"- {path}" for path in ordered[:10])
            raise RepositoryError(
                f"Foram encontradas {len(ordered)} cópias de '{candidate_name}'. "
                f"Configure EXCEL_SEARCH_ROOTS ou EXCEL_PATH para indicar a correta:\n{locations}"
            )

    if _sample_fallback_enabled():
        for candidate_name in filenames:
            sample = (base_dir / "data" / candidate_name).resolve()
            if sample.is_file():
                log.warning("Arquivo sincronizado não encontrado; usando a base de demonstração: %s", sample)
                return sample

    searched = ", ".join(str(root) for root in roots) or "nenhuma pasta OneDrive detectada"
    requested = ", ".join(f"'{name}'" for name in filenames)
    raise RepositoryError(
        f"Nenhum dos arquivos {requested} foi encontrado nas pastas OneDrive/SharePoint sincronizadas ({searched})."
    )


class LocalExcelRepository:
    def __init__(self, base_dir: Path):
        self.base_dir = base_dir
        # EXCEL_FILE_NAME was used by the first release. Keep it as a fallback
        # so an existing .env continues to work after upgrading.
        self.file_name = (
            os.getenv("EXCEL_FILENAME", "").strip()
            or os.getenv("EXCEL_FILE_NAME", "").strip()
            or "B2B_CTACUSTOS.xlsx"
        )
        self.sheet_name = os.getenv("EXCEL_SHEET_NAME", "Atividades")
        self._lock = threading.RLock()
        self.file_path = self._locate_file()
        sample_path = (self.base_dir / "data" / self.file_name).resolve()
        self.is_sample = os.path.normcase(str(self.file_path)) == os.path.normcase(str(sample_path))

    @property
    def source_label(self) -> str:
        source_type = "Excel de demonstração" if self.is_sample else "Excel sincronizado"
        return f"{source_type} · {self.file_path.name} · aba {self.sheet_name}"

    def _locate_file(self) -> Path:
        explicit = os.getenv("EXCEL_PATH", "").strip()
        if explicit:
            path = Path(explicit).expanduser()
            if not path.is_absolute():
                path = self.base_dir / path
            path = path.resolve()
            if not path.is_file():
                raise RepositoryError(f"Arquivo Excel não encontrado: {path}")
            if path.suffix.casefold() not in {".xlsx", ".xlsm"}:
                raise RepositoryError("EXCEL_PATH deve apontar para um arquivo .xlsx ou .xlsm.")
            return path
        aliases = [
            item.strip()
            for item in os.getenv(
                "EXCEL_FILENAME_ALIASES",
                "",
            ).split(";")
            if item.strip()
        ]
        return discover_excel_file(
            self.file_name,
            base_dir=self.base_dir,
            fallback_filenames=aliases,
        )

    def _openpyxl(self):
        try:
            import openpyxl
        except ImportError as exc:
            raise RepositoryError("Instale a dependência openpyxl: pip install -r requirements.txt") from exc
        return openpyxl

    def _load_workbook(self, openpyxl, **options):
        return self._load_workbook_path(openpyxl, self.file_path, **options)

    @staticmethod
    def _load_workbook_path(openpyxl, path: Path, **options):
        try:
            return openpyxl.load_workbook(path, **options)
        except PermissionError:
            log.info("Excel aberto em outro processo; usando leitura compartilhada: %s", path)
            workbook_bytes = _read_file_with_shared_access(path)
            return openpyxl.load_workbook(io.BytesIO(workbook_bytes), **options)

    def _locate_reference_workbook(self, kind: str) -> Path:
        definitions = {
            "services": (
                "SERVICES",
                "SERVICOS.xlsx",
                ("SERVIÇOS.xlsx", "SERVICOS.xlsm", "SERVIÇOS.xlsm"),
                "serviços",
            ),
            "materials": (
                "MATERIALS",
                "MATERIAL.xlsx",
                ("MATERIAIS.xlsx", "MATERIAL.xlsm", "MATERIAIS.xlsm"),
                "materiais",
            ),
        }
        if kind not in definitions:
            raise ValueError("Tipo de planilha de referência inválido.")
        prefix, default_name, default_aliases, label = definitions[kind]

        explicit = os.getenv(f"{prefix}_EXCEL_PATH", "").strip()
        if explicit:
            path = Path(explicit).expanduser()
            if not path.is_absolute():
                path = self.file_path.parent / path
            path = path.resolve()
            if path.is_file() and path.suffix.casefold() in {".xlsx", ".xlsm"}:
                return path
            if path.is_file():
                raise RepositoryError(f"A planilha de {label} deve ser .xlsx ou .xlsm.")
            log.warning(
                "%s_EXCEL_PATH aponta para um arquivo que não existe neste computador; "
                "tentando a descoberta automática: %s",
                prefix,
                path,
            )

        try:
            metadata = self.get_settings().get("uploads", {}).get(kind, {})
            uploaded_path = Path(clean_text(metadata.get("path"), 500)).expanduser()
            if str(uploaded_path) not in {"", "."} and uploaded_path.is_file():
                return uploaded_path.resolve()
        except (OSError, RepositoryError):
            pass

        filename = os.getenv(f"{prefix}_EXCEL_FILENAME", default_name).strip() or default_name
        configured_aliases = [
            item.strip()
            for item in os.getenv(f"{prefix}_EXCEL_FILENAME_ALIASES", "").split(";")
            if item.strip()
        ]
        candidate_names = [filename, *configured_aliases, *default_aliases]
        unique_names: list[str] = []
        seen_names: set[str] = set()
        for candidate_name in candidate_names:
            normalized_name = candidate_name.casefold()
            if normalized_name not in seen_names:
                seen_names.add(normalized_name)
                unique_names.append(candidate_name)

        for directory in (self.file_path.parent, self._reference_upload_dir()):
            for candidate_name in unique_names:
                try:
                    match = next(
                        (
                            child
                            for child in directory.iterdir()
                            if child.is_file() and child.name.casefold() == candidate_name.casefold()
                        ),
                        None,
                    )
                except (FileNotFoundError, OSError, PermissionError):
                    match = None
                if match:
                    return match.resolve()

        try:
            return discover_excel_file(
                filename,
                base_dir=self.base_dir,
                fallback_filenames=unique_names[1:],
                include_automatic_roots=True,
            )
        except RepositoryError as exc:
            accepted_names = ", ".join(unique_names)
            raise RepositoryError(
                f"Não foi possível localizar a planilha de {label}. Nomes procurados: {accepted_names}. {exc} "
                "No outro computador, confirme no Explorador de Arquivos que a biblioteca do SharePoint "
                "está sincronizada pelo OneDrive e que o arquivo está disponível localmente. "
                f"Como alternativa, configure {prefix}_EXCEL_PATH no .env."
            ) from exc

    @staticmethod
    def _reference_sheet(workbook, kind: str):
        definitions = {
            "services": ("SERVICES", "SERVICOS", "serviços"),
            "materials": ("MATERIALS", "MATERIAIS", "materiais"),
        }
        prefix, default_sheet, label = definitions[kind]
        configured = os.getenv(f"{prefix}_EXCEL_SHEET_NAME", "").strip()
        requested = configured or default_sheet
        normalized_requested = normalize_header(requested)
        matching_sheet = next(
            (name for name in workbook.sheetnames if normalize_header(name) == normalized_requested),
            None,
        )
        if matching_sheet:
            return workbook[matching_sheet]
        if len(workbook.sheetnames) == 1:
            if configured:
                log.warning(
                    "A aba '%s' não existe em %s; usando a única aba disponível: %s",
                    requested,
                    label,
                    workbook.sheetnames[0],
                )
            return workbook[workbook.sheetnames[0]]
        raise RepositoryError(
            f"Aba '{requested}' não encontrada na planilha de {label}. "
            f"Abas disponíveis: {', '.join(workbook.sheetnames)}"
        )

    def list_services(self) -> dict[str, Any]:
        path = self._locate_reference_workbook("services")
        openpyxl = self._openpyxl()
        try:
            workbook = self._load_workbook_path(
                openpyxl,
                path,
                read_only=True,
                data_only=True,
                keep_vba=path.suffix.casefold() == ".xlsm",
            )
        except (OSError, ValueError) as exc:
            raise RepositoryError(f"Não foi possível abrir a planilha de serviços: {exc}") from exc

        try:
            sheet = self._reference_sheet(workbook, "services")
            header_row = None
            columns: dict[str, int] = {}
            for row_number in range(1, min(sheet.max_row, 10) + 1):
                candidate = resolve_service_columns(
                    sheet.cell(row_number, column).value
                    for column in range(1, sheet.max_column + 1)
                )
                if {"description", "unit_price"}.issubset(candidate):
                    header_row = row_number
                    columns = candidate
                    break
            if header_row is None:
                raise RepositoryError(
                    "A planilha de serviços precisa das colunas SERVICO e CUSTO_UNITARIO."
                )

            services: list[dict[str, Any]] = []
            used_keys: set[str] = set()
            for row_number in range(header_row + 1, sheet.max_row + 1):
                description = clean_text(
                    sheet.cell(row_number, columns["description"] + 1).value,
                    500,
                )
                if not description:
                    continue
                unit_price = round(
                    to_number(sheet.cell(row_number, columns["unit_price"] + 1).value),
                    4,
                )
                if unit_price < 0:
                    raise RepositoryError(
                        f"O serviço da linha {row_number} possui valor unitário negativo."
                    )
                code = (
                    clean_text(sheet.cell(row_number, columns["code"] + 1).value, 120)
                    if "code" in columns
                    else ""
                )
                unit = (
                    clean_text(sheet.cell(row_number, columns["unit"] + 1).value, 40)
                    if "unit" in columns
                    else ""
                )
                identity = "|".join((normalize_header(code), normalize_header(description), normalize_header(unit)))
                service_key = "svc-" + hashlib.sha1(identity.encode("utf-8")).hexdigest()[:16]
                if service_key in used_keys:
                    service_key = f"{service_key}-{row_number}"
                used_keys.add(service_key)
                services.append(
                    {
                        "key": service_key,
                        "code": code,
                        "description": description,
                        "unit": unit,
                        "unit_price": unit_price,
                    }
                )
            return {
                "items": services,
                "source": f"{path.name} · aba {sheet.title}",
            }
        finally:
            workbook.close()

    def list_materials(self) -> dict[str, Any]:
        path = self._locate_reference_workbook("materials")
        openpyxl = self._openpyxl()
        try:
            workbook = self._load_workbook_path(
                openpyxl,
                path,
                read_only=True,
                data_only=True,
                keep_vba=path.suffix.casefold() == ".xlsm",
            )
        except (OSError, ValueError) as exc:
            raise RepositoryError(f"Não foi possível abrir a planilha de materiais: {exc}") from exc

        try:
            sheet = self._reference_sheet(workbook, "materials")
            header_row = None
            columns: dict[str, int] = {}
            for row_number in range(1, min(sheet.max_row, 10) + 1):
                candidate = resolve_material_columns(
                    sheet.cell(row_number, column).value
                    for column in range(1, sheet.max_column + 1)
                )
                if {"description", "unit_price"}.issubset(candidate):
                    header_row = row_number
                    columns = candidate
                    break
            if header_row is None:
                raise RepositoryError(
                    "A planilha de materiais precisa das colunas Material e Valor."
                )

            materials: list[dict[str, Any]] = []
            used_keys: set[str] = set()
            for row_number in range(header_row + 1, sheet.max_row + 1):
                description = clean_text(
                    sheet.cell(row_number, columns["description"] + 1).value,
                    500,
                )
                if not description:
                    continue
                unit_price = round(
                    to_number(sheet.cell(row_number, columns["unit_price"] + 1).value),
                    4,
                )
                if unit_price < 0:
                    raise RepositoryError(
                        f"O material da linha {row_number} possui valor unitário negativo."
                    )
                code = (
                    clean_text(sheet.cell(row_number, columns["code"] + 1).value, 120)
                    if "code" in columns
                    else ""
                )
                unit = (
                    clean_text(sheet.cell(row_number, columns["unit"] + 1).value, 40)
                    if "unit" in columns
                    else ""
                )
                identity = "|".join((normalize_header(code), normalize_header(description), normalize_header(unit)))
                material_key = "mat-" + hashlib.sha1(identity.encode("utf-8")).hexdigest()[:16]
                if material_key in used_keys:
                    material_key = f"{material_key}-{row_number}"
                used_keys.add(material_key)
                materials.append(
                    {
                        "key": material_key,
                        "code": code,
                        "description": description,
                        "unit": unit,
                        "unit_price": unit_price,
                    }
                )
            return {
                "items": materials,
                "source": f"{path.name} · aba {sheet.title}",
            }
        finally:
            workbook.close()

    def list_activities(self) -> list[dict[str, Any]]:
        with self._lock:
            openpyxl = self._openpyxl()
            try:
                workbook = self._load_workbook(
                    openpyxl,
                    read_only=True,
                    data_only=False,
                    keep_vba=self.file_path.suffix.casefold() == ".xlsm",
                )
            except (OSError, ValueError) as exc:
                raise RepositoryError(f"Não foi possível abrir '{self.file_path.name}': {exc}") from exc
            try:
                if self.sheet_name not in workbook.sheetnames:
                    raise RepositoryError(
                        f"Aba '{self.sheet_name}' não encontrada. Abas disponíveis: {', '.join(workbook.sheetnames)}"
                    )
                sheet = workbook[self.sheet_name]
                rows = sheet.iter_rows(values_only=True)
                headers = next(rows, None)
                if not headers:
                    return []
                field_columns = resolve_field_columns(headers)
                required_fields = ("id", "tipo_atividade", "status", "situacao")
                missing = [HEADERS[FIELD_NAMES.index(field)] for field in required_fields if field not in field_columns]
                if missing:
                    raise RepositoryError(
                        "Colunas obrigatórias ausentes no Excel: "
                        f"{', '.join(missing)}. Também são aceitos Obra Executada, Status Obra e Detalhes Obra."
                    )

                default_daily_rate = self._daily_rate_from_open_workbook(workbook)
                activities = []
                for row_number, row in enumerate(rows, start=2):
                    if not row or all(value in (None, "") for value in row):
                        continue
                    raw = {}
                    for field in FIELD_NAMES:
                        index = field_columns.get(field)
                        raw[field] = row[index] if index is not None and index < len(row) else None
                    if raw.get("id") in (None, ""):
                        continue
                    raw["calculation_mode"] = calculation_mode(field_columns)
                    if to_number(raw.get("custo_tecnico_dia")) <= 0:
                        raw["custo_tecnico_dia"] = default_daily_rate
                    activity = normalize_activity(raw)
                    activity["record_key"] = self._record_key(row_number, activity["id"])
                    activities.append(activity)
                return activities
            finally:
                workbook.close()

    def _activity_options(self, activities: list[dict[str, Any]]) -> dict[str, Any]:
        def unique(field: str) -> list[str]:
            return sorted(
                {clean_text(item.get(field)) for item in activities if clean_text(item.get(field))},
                key=str.casefold,
            )

        technicians: dict[str, dict[str, str]] = {}
        for item in activities:
            name = clean_text(item.get("tecnico_nome"))
            registration = clean_text(item.get("matricula"))
            if not name and not registration:
                continue
            key = f"{registration}|{name}".casefold()
            technicians[key] = {"name": name, "registration": registration}

        return {
            "statuses": unique("status"),
            "types": unique("tipo_atividade"),
            "companies": unique("empresa"),
            "eps": unique("eps"),
            "regions": unique("regiao"),
            "clusters": unique("cluster"),
            "technologies": unique("tecnologia"),
            "technicians": sorted(
                technicians.values(),
                key=lambda item: (item["name"].casefold(), item["registration"].casefold()),
            ),
        }

    def _reference_upload_dir(self) -> Path:
        configured = os.getenv("REFERENCE_UPLOAD_DIR", "").strip()
        if configured:
            directory = Path(configured).expanduser()
            if not directory.is_absolute():
                directory = self.file_path.parent / directory
        else:
            directory = self.file_path.parent / "B2B_CTACUSTOS_Importacoes"
        return directory.resolve()

    @staticmethod
    def _config_header_row(sheet) -> int | None:
        for row_number in range(4, min(sheet.max_row, 12) + 1):
            if normalize_header(sheet.cell(row_number, 1).value) == normalize_header("Categoria"):
                return row_number
        return None

    @staticmethod
    def _daily_rate_from_open_workbook(workbook) -> float:
        """Read the configured daily rate without reopening the workbook."""
        reference_month = date.today().strftime("%Y-%m")
        business_days = business_days_for_month(reference_month)
        fallback_monthly_cost = 15_000.0
        fallback_rate = round(fallback_monthly_cost / business_days, 2)
        if "Config" not in workbook.sheetnames:
            return fallback_rate

        sheet = workbook["Config"]
        is_current_layout = normalize_header(sheet["A2"].value) == normalize_header("Custo mensal técnico")
        if not is_current_layout:
            legacy_daily_rate = to_number(sheet["B2"].value)
            return legacy_daily_rate if legacy_daily_rate > 0 else fallback_rate

        saved_month = clean_text(sheet["B5"].value, 7)
        try:
            datetime.strptime(saved_month, "%Y-%m")
            reference_month = saved_month
        except ValueError:
            pass
        saved_days = int(to_number(sheet["B3"].value))
        business_days = saved_days if saved_days > 0 else business_days_for_month(reference_month)
        monthly_cost = max(0.0, to_number(sheet["B2"].value))
        return round(monthly_cost / business_days, 2) if business_days > 0 else 0.0

    def _read_settings(self, activity_options: dict[str, Any]) -> dict[str, Any]:
        reference_month = date.today().strftime("%Y-%m")
        business_days = business_days_for_month(reference_month)
        monthly_cost = 15_000.0
        technology_names: dict[str, str] = {}
        for name in (*activity_options["technologies"], *DEFAULT_TECHNOLOGIES):
            text = clean_text(name, 160)
            if text:
                technology_names.setdefault(text.casefold(), text)
        technology_rates = [
            {"name": name, "service_cost": 0.0, "configured": False}
            for name in technology_names.values()
        ]
        settings = {
            "reference_month": reference_month,
            "monthly_technician_cost": monthly_cost,
            "business_days": business_days,
            "global_daily_rate": round(monthly_cost / business_days, 2),
            "statuses": list(activity_options["statuses"]),
            "types": list(activity_options["types"]),
            "technologies": [item["name"] for item in technology_rates],
            "technology_rates": technology_rates,
            "category_teams": [
                {"category": category["key"], "label": category["label"], "technicians": 0}
                for category in SERVICE_CATEGORIES
            ],
            "technicians": [dict(item) for item in activity_options["technicians"]],
            "companies": list(activity_options["companies"]),
            "eps": list(activity_options["eps"]),
            "upload_directory": str(self._reference_upload_dir()),
            "uploads": {
                "services": {"filename": "", "path": "", "uploaded_at": ""},
                "materials": {"filename": "", "path": "", "uploaded_at": ""},
            },
            "calculation_status": (
                "Valor da equipe por categoria = total de técnicos × custo mensal do técnico. "
                "O custo diário permanece disponível somente para a calculadora detalhada."
            ),
        }
        openpyxl = self._openpyxl()
        try:
            workbook = self._load_workbook(
                openpyxl,
                read_only=True,
                data_only=True,
                keep_vba=self.file_path.suffix.casefold() == ".xlsm",
            )
            try:
                if "Config" not in workbook.sheetnames:
                    return settings
                sheet = workbook["Config"]
                is_current_layout = normalize_header(sheet["A2"].value) == normalize_header("Custo mensal técnico")
                if is_current_layout:
                    saved_month = clean_text(sheet["B5"].value, 7)
                    try:
                        datetime.strptime(saved_month, "%Y-%m")
                        settings["reference_month"] = saved_month
                    except ValueError:
                        pass
                    settings["monthly_technician_cost"] = max(0.0, to_number(sheet["B2"].value))
                    saved_days = int(to_number(sheet["B3"].value))
                    settings["business_days"] = saved_days if saved_days > 0 else business_days_for_month(settings["reference_month"])
                    settings["global_daily_rate"] = round(
                        settings["monthly_technician_cost"] / settings["business_days"], 2
                    )
                else:
                    legacy_daily_rate = to_number(sheet["B2"].value)
                    if legacy_daily_rate > 0:
                        settings["global_daily_rate"] = legacy_daily_rate
                        settings["monthly_technician_cost"] = round(legacy_daily_rate * business_days, 2)

                # A Config sheet with a Categoria header is authoritative. Older
                # layouts without it still bootstrap choices from activity rows.
                header_row = self._config_header_row(sheet)
                if header_row is None:
                    return settings

                configured = {
                    "statuses": [],
                    "types": [],
                    "technologies": [],
                    "technology_rates": [],
                    "technicians": [],
                    "companies": [],
                    "eps": [],
                    "category_teams": [],
                }
                configured_catalogs: set[str] = set()
                technology_catalog_has_rates = False
                for row in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                    category = normalize_header(row[0] if len(row) > 0 else "")
                    raw_code = row[1] if len(row) > 1 else ""
                    code = clean_text(raw_code, 120)
                    value = clean_text(row[2] if len(row) > 2 else "", 500)
                    uploaded_at = clean_text(row[3] if len(row) > 3 else "", 40)
                    if category == "status" and value:
                        configured["statuses"].append(value)
                    elif category in {"tipoatividade", "tiposatividade"}:
                        configured_catalogs.add("types")
                        if value:
                            configured["types"].append(value)
                    elif category in {"tecnologia", "tecnologias"}:
                        configured_catalogs.add("technologies")
                        if value:
                            configured["technologies"].append(value)
                            configured["technology_rates"].append(
                                {
                                    "name": value,
                                    "service_cost": max(0.0, to_number(raw_code)),
                                    "configured": raw_code not in (None, ""),
                                }
                            )
                            technology_catalog_has_rates = technology_catalog_has_rates or raw_code not in (None, "")
                    elif category in {"equipecategoria", "categoriaequipe"} and code:
                        category_key = service_category(code) or normalize_header(code)
                        if category_key in {item["key"] for item in SERVICE_CATEGORIES}:
                            configured["category_teams"].append(
                                {
                                    "category": category_key,
                                    "label": next(
                                        item["label"] for item in SERVICE_CATEGORIES if item["key"] == category_key
                                    ),
                                    "technicians": max(0, int(to_number(value))),
                                }
                            )
                    elif category in {"tecnico", "tecnicos"} and (code or value):
                        configured["technicians"].append({"registration": code, "name": value})
                    elif category in {"empresa", "empresas"} and value:
                        configured["companies"].append(value)
                    elif category == "eps" and value:
                        configured["eps"].append(value)
                    elif category == "arquivoservicos" and (code or value):
                        settings["uploads"]["services"] = {
                            "filename": code,
                            "path": value,
                            "uploaded_at": uploaded_at,
                        }
                    elif category == "arquivomateriais" and (code or value):
                        settings["uploads"]["materials"] = {
                            "filename": code,
                            "path": value,
                            "uploaded_at": uploaded_at,
                        }

                # Older Config sheets predate these two catalogs. Bootstrap them
                # from the activity base until the administrator saves the new layout.
                if "types" not in configured_catalogs:
                    configured["types"] = settings["types"]
                if "technologies" not in configured_catalogs:
                    configured["technologies"] = settings["technologies"]
                    configured["technology_rates"] = settings["technology_rates"]
                elif not technology_catalog_has_rates:
                    # Migração transparente do cadastro antigo, que possuía
                    # somente o nome da tecnologia e nenhuma coluna de valor.
                    merged_rates: dict[str, dict[str, Any]] = {
                        item["name"].casefold(): item for item in configured["technology_rates"]
                    }
                    for item in settings["technology_rates"]:
                        merged_rates.setdefault(item["name"].casefold(), item)
                    configured["technology_rates"] = list(merged_rates.values())
                    configured["technologies"] = [item["name"] for item in configured["technology_rates"]]
                configured_team_map = {
                    item["category"]: item for item in configured["category_teams"]
                }
                configured["category_teams"] = [
                    configured_team_map.get(
                        category["key"],
                        {"category": category["key"], "label": category["label"], "technicians": 0},
                    )
                    for category in SERVICE_CATEGORIES
                ]
                settings.update(configured)
                return settings
            finally:
                workbook.close()
        except (OSError, ValueError) as exc:
            raise RepositoryError(f"Não foi possível ler as configurações: {exc}") from exc

    def get_options(self) -> dict[str, Any]:
        activity_options = self._activity_options(self.list_activities())
        settings = self._read_settings(activity_options)
        activity_options.update(
            {
                "statuses": settings["statuses"],
                "types": settings["types"],
                "technologies": settings["technologies"],
                "technology_rates": settings["technology_rates"],
                "category_teams": settings["category_teams"],
                "monthly_technician_cost": settings["monthly_technician_cost"],
                "technicians": settings["technicians"],
                "companies": settings["companies"],
                "eps": settings["eps"],
                "default_daily_rate": settings["global_daily_rate"],
            }
        )
        return activity_options

    def _ensure_draft_column(self, sheet) -> None:
        """Append DRAFT without disturbing the existing table formatting."""
        if "draft" in resolve_field_columns(cell.value for cell in sheet[1]):
            return

        from openpyxl.utils import get_column_letter, range_boundaries
        from openpyxl.worksheet.table import TableColumn

        previous_column = max(1, sheet.max_column)
        draft_column = previous_column + 1
        header_source = sheet.cell(1, previous_column)
        header_target = sheet.cell(1, draft_column, "DRAFT")
        if header_source.has_style:
            header_target._style = copy(header_source._style)
        header_target.number_format = header_source.number_format

        for row_number in range(2, sheet.max_row + 1):
            source = sheet.cell(row_number, previous_column)
            target = sheet.cell(row_number, draft_column)
            if source.has_style:
                target._style = copy(source._style)
            target.number_format = source.number_format

        previous_letter = get_column_letter(previous_column)
        draft_letter = get_column_letter(draft_column)
        previous_dimension = sheet.column_dimensions[previous_letter]
        draft_dimension = sheet.column_dimensions[draft_letter]
        draft_dimension.width = previous_dimension.width or 18
        draft_dimension.hidden = previous_dimension.hidden

        for table in sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if min_row == 1 and max_col == previous_column:
                table.ref = f"{get_column_letter(min_col)}{min_row}:{draft_letter}{max_row}"
                if table.tableColumns:
                    next_id = max(column.id for column in table.tableColumns) + 1
                    table.tableColumns.append(TableColumn(id=next_id, name="DRAFT"))
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref

        if sheet.auto_filter.ref:
            min_col, min_row, max_col, max_row = range_boundaries(sheet.auto_filter.ref)
            if min_row == 1 and max_col == previous_column:
                sheet.auto_filter.ref = f"{get_column_letter(min_col)}{min_row}:{draft_letter}{max_row}"

    def _schema(self, sheet) -> tuple[list[Any], dict[str, int], str]:
        headers = [cell.value for cell in sheet[1]]
        zero_based = resolve_field_columns(headers)
        columns = {field: index + 1 for field, index in zero_based.items()}
        if "id" not in columns:
            raise RepositoryError("Coluna ID não encontrada no Excel.")
        return headers, columns, calculation_mode(zero_based)

    @staticmethod
    def _record_key(row_number: int, activity_id: Any) -> str:
        return f"{row_number}:{clean_text(activity_id, 60)}"

    def _find_activity_row(
        self,
        sheet,
        field_columns: dict[str, int],
        activity_reference: str,
    ) -> int:
        reference = clean_text(activity_reference, 200)
        if not reference:
            raise ValueError("Atividade não informada.")

        row_match = re.fullmatch(r"(\d+):(.*)", reference)
        if row_match:
            row_number = int(row_match.group(1))
            expected_id = clean_text(row_match.group(2), 60)
            if 2 <= row_number <= sheet.max_row:
                actual_id = clean_text(sheet.cell(row_number, field_columns["id"]).value, 60)
                if actual_id == expected_id:
                    return row_number
            raise RepositoryError(
                "A linha da atividade mudou no Excel. Atualize a lista e abra a calculadora novamente."
            )

        matching_rows = [
            row_number
            for row_number in range(2, sheet.max_row + 1)
            if clean_text(sheet.cell(row_number, field_columns["id"]).value, 60) == reference
        ]
        if not matching_rows:
            raise RepositoryError(f"Atividade '{reference}' não encontrada.")
        if len(matching_rows) > 1:
            raise RepositoryError(
                f"O ID '{reference}' aparece em mais de uma linha. Atualize a lista para selecionar a atividade exata."
            )
        return matching_rows[0]

    @staticmethod
    def _row_payload(sheet, field_columns: dict[str, int], row_number: int) -> dict[str, Any]:
        return {
            field: sheet.cell(row_number, column).value
            for field, column in field_columns.items()
            if field != "id"
        }

    def _prepare_activity(
        self,
        activity_id: str,
        payload: dict[str, Any],
        mode: str,
        *,
        use_team_gap: bool = True,
        technology_rates: list[dict[str, Any]] | None = None,
    ) -> dict[str, Any]:
        prepared_payload = dict(payload)
        technology_key = normalize_header(prepared_payload.get("tecnologia"))
        if technology_key and technology_rates is not None:
            rate_by_technology = {
                normalize_header(item.get("name")): max(0.0, to_number(item.get("service_cost")))
                for item in technology_rates
                if (
                    isinstance(item, dict)
                    and clean_text(item.get("name"))
                    and item.get("configured", True) is not False
                )
            }
            if technology_key in rate_by_technology:
                prepared_payload["custo_mo"] = rate_by_technology[technology_key]
        activity = validate_update(activity_id, prepared_payload, use_team_gap=use_team_gap)
        activity["calculation_mode"] = mode
        activity["_use_team_gap"] = use_team_gap
        if not activity.get("data"):
            activity["data"] = date.today().isoformat()
        # Regra de negócio: o custo evitado representa o valor da mão de obra.
        activity["custo_evitado"] = round(to_number(activity.get("custo_mo")), 2)
        activity.update(calculate_values(activity, use_team_gap=use_team_gap))
        return activity

    def _apply_activity_to_row(
        self,
        sheet,
        headers: list[Any],
        field_columns: dict[str, int],
        target_row: int,
        activity: dict[str, Any],
        mode: str,
        *,
        include_id: bool,
    ) -> None:
        from openpyxl.utils import get_column_letter

        for field in FIELD_NAMES:
            column = field_columns.get(field)
            if not column or (field == "id" and not include_id):
                continue

            cell = sheet.cell(target_row, column)
            if field == "id":
                identifier = clean_text(activity.get("id"), 60)
                cell.value = int(identifier) if identifier.isdigit() else identifier
            elif field == "custo_total":
                if mode == "detailed":
                    material_cost = get_column_letter(field_columns["custo_mat"])
                    labor_cost = get_column_letter(field_columns["custo_mo"])
                    cell.value = f"={material_cost}{target_row}+{labor_cost}{target_row}"
                else:
                    cell.value = activity["custo_total"]
            elif field == "custo_por_tecnico":
                if "custo_total" in field_columns and "qtde_tecnicos" in field_columns:
                    total = get_column_letter(field_columns["custo_total"])
                    technicians = get_column_letter(field_columns["qtde_tecnicos"])
                    cell.value = f"=IF({technicians}{target_row}>0,{total}{target_row}/{technicians}{target_row},0)"
                else:
                    cell.value = activity["custo_por_tecnico"]
            elif field == "gap":
                gap_dependencies = ("custo_mo", "custo_tecnico_dia", "qtde_tecnicos", "qtde_dias")
                if activity.get("_use_team_gap") and all(dependency in field_columns for dependency in gap_dependencies):
                    labor = get_column_letter(field_columns["custo_mo"])
                    daily_rate = get_column_letter(field_columns["custo_tecnico_dia"])
                    technicians = get_column_letter(field_columns["qtde_tecnicos"])
                    days = get_column_letter(field_columns["qtde_dias"])
                    cell.value = (
                        f"={labor}{target_row}-"
                        f"({daily_rate}{target_row}*{technicians}{target_row}*{days}{target_row})"
                    )
                else:
                    cell.value = activity["gap"]
            elif field == "atualizado_em":
                if normalize_header(headers[column - 1]) == normalize_header("Atualizado Em"):
                    cell.value = datetime.now()
            elif field == "data":
                if normalize_header(headers[column - 1]) == normalize_header("Data"):
                    parsed = activity_date(activity.get("data"))
                    cell.value = datetime.combine(parsed, datetime.min.time()) if parsed else None
            else:
                cell.value = activity.get(field)

    def _copy_row_formatting(self, sheet, source_row: int, target_row: int) -> None:
        from openpyxl.utils import get_column_letter

        if source_row < 2:
            return
        for column in range(1, sheet.max_column + 1):
            source = sheet.cell(source_row, column)
            target = sheet.cell(target_row, column)
            if source.has_style:
                target._style = copy(source._style)
            if source.number_format:
                target.number_format = source.number_format
        source_dimension = sheet.row_dimensions[source_row]
        target_dimension = sheet.row_dimensions[target_row]
        target_dimension.height = source_dimension.height
        target_dimension.hidden = source_dimension.hidden

        validations = getattr(sheet, "data_validations", None)
        if validations:
            for validation in validations.dataValidation:
                for cell_range in list(validation.sqref.ranges):
                    if sheet.cell(source_row, cell_range.min_col).coordinate in cell_range:
                        validation.add(
                            f"{get_column_letter(cell_range.min_col)}{target_row}:"
                            f"{get_column_letter(cell_range.max_col)}{target_row}"
                        )

    def _extend_tables(self, sheet, previous_row: int, target_row: int) -> None:
        from openpyxl.utils import get_column_letter, range_boundaries

        for table in sheet.tables.values():
            min_col, min_row, max_col, max_row = range_boundaries(table.ref)
            if max_row == previous_row and min_row < target_row:
                table.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{target_row}"
                )
                if table.autoFilter is not None:
                    table.autoFilter.ref = table.ref
        if sheet.auto_filter.ref:
            min_col, min_row, max_col, max_row = range_boundaries(sheet.auto_filter.ref)
            if max_row == previous_row:
                sheet.auto_filter.ref = (
                    f"{get_column_letter(min_col)}{min_row}:"
                    f"{get_column_letter(max_col)}{target_row}"
                )

    def _save_workbook(self, workbook) -> None:
        if getattr(workbook, "calculation", None) is not None:
            workbook.calculation.fullCalcOnLoad = True
            workbook.calculation.forceFullCalc = True
        temp_path = self.file_path.with_name(f"{self.file_path.stem}.saving{self.file_path.suffix}")
        try:
            workbook.save(temp_path)
            workbook.close()
            os.replace(temp_path, self.file_path)
        finally:
            if temp_path.exists():
                try:
                    temp_path.unlink()
                except OSError:
                    pass

    def update_activity(self, activity_id: str, payload: dict[str, Any]) -> dict[str, Any]:
        settings = self.get_settings()
        with self._lock:
            openpyxl = self._openpyxl()
            try:
                workbook = self._load_workbook(openpyxl, keep_vba=self.file_path.suffix.casefold() == ".xlsm")
                if self.sheet_name not in workbook.sheetnames:
                    raise RepositoryError(f"Aba '{self.sheet_name}' não encontrada.")
                sheet = workbook[self.sheet_name]
                self._ensure_draft_column(sheet)
                headers, field_columns, mode = self._schema(sheet)

                target_row = self._find_activity_row(sheet, field_columns, activity_id)

                # Fields intentionally hidden from the modal must keep their
                # historical values when another field is edited.
                merged_payload = self._row_payload(sheet, field_columns, target_row)
                merged_payload.update(payload)
                merged_payload["custo_tecnico_dia"] = self._daily_rate_from_open_workbook(workbook)
                actual_id = clean_text(sheet.cell(target_row, field_columns["id"]).value, 60)
                activity = self._prepare_activity(
                    actual_id,
                    merged_payload,
                    mode,
                    use_team_gap=False,
                    technology_rates=settings.get("technology_rates", []),
                )

                self._apply_activity_to_row(
                    sheet, headers, field_columns, target_row, activity, mode, include_id=False
                )
                self._save_workbook(workbook)
                saved = normalize_activity(activity, use_team_gap=False)
                saved["record_key"] = self._record_key(target_row, saved["id"])
                return saved
            except PermissionError as exc:
                raise RepositoryError("O Excel está bloqueado para edição. Feche o arquivo ou verifique a sincronização.") from exc
            except OSError as exc:
                raise RepositoryError(f"Não foi possível salvar o Excel: {exc}") from exc

    def create_activity(self, payload: dict[str, Any]) -> dict[str, Any]:
        settings = self.get_settings()
        with self._lock:
            openpyxl = self._openpyxl()
            try:
                workbook = self._load_workbook(openpyxl, keep_vba=self.file_path.suffix.casefold() == ".xlsm")
                if self.sheet_name not in workbook.sheetnames:
                    raise RepositoryError(f"Aba '{self.sheet_name}' não encontrada.")
                sheet = workbook[self.sheet_name]
                self._ensure_draft_column(sheet)
                headers, field_columns, mode = self._schema(sheet)

                existing_ids = {
                    clean_text(sheet.cell(row, field_columns["id"]).value)
                    for row in range(2, sheet.max_row + 1)
                    if clean_text(sheet.cell(row, field_columns["id"]).value)
                }
                requested_id = clean_text(payload.get("id"), 60)
                if requested_id and requested_id in existing_ids:
                    raise RepositoryError(f"Já existe uma atividade com o ID '{requested_id}'.")
                numeric_ids = [int(value) for value in existing_ids if value.isdigit()]
                activity_id = requested_id or str(max(numeric_ids, default=0) + 1)
                prepared_payload = dict(payload)
                prepared_payload["custo_tecnico_dia"] = self._daily_rate_from_open_workbook(workbook)
                activity = self._prepare_activity(
                    activity_id,
                    prepared_payload,
                    mode,
                    use_team_gap=False,
                    technology_rates=settings.get("technology_rates", []),
                )

                target_row = sheet.max_row + 1
                previous_row = target_row - 1
                self._copy_row_formatting(sheet, previous_row, target_row)
                self._apply_activity_to_row(
                    sheet, headers, field_columns, target_row, activity, mode, include_id=True
                )
                self._extend_tables(sheet, previous_row, target_row)
                self._save_workbook(workbook)
                saved = normalize_activity(activity, use_team_gap=False)
                saved["record_key"] = self._record_key(target_row, saved["id"])
                return saved
            except PermissionError as exc:
                raise RepositoryError("O Excel está bloqueado para inclusão. Feche o arquivo ou verifique a sincronização.") from exc
            except OSError as exc:
                raise RepositoryError(f"Não foi possível incluir no Excel: {exc}") from exc

    def export_activities(self) -> tuple[str, bytes]:
        """Export the activity sheet with its formatting and without DRAFT."""
        with self._lock:
            openpyxl = self._openpyxl()
            try:
                workbook = self._load_workbook(
                    openpyxl,
                    data_only=False,
                    keep_vba=self.file_path.suffix.casefold() == ".xlsm",
                )
            except (OSError, ValueError) as exc:
                raise RepositoryError(f"Não foi possível abrir a base para exportação: {exc}") from exc

            try:
                if self.sheet_name not in workbook.sheetnames:
                    raise RepositoryError(f"Aba '{self.sheet_name}' não encontrada.")
                source = workbook[self.sheet_name]
                field_columns = resolve_field_columns(cell.value for cell in source[1])
                excluded_column = field_columns.get("draft")
                excluded_column = excluded_column + 1 if excluded_column is not None else None

                exported = openpyxl.Workbook()
                target = exported.active
                target.title = source.title[:31]
                target.sheet_view.showGridLines = source.sheet_view.showGridLines
                target.sheet_format.defaultRowHeight = source.sheet_format.defaultRowHeight
                if source.sheet_properties.tabColor:
                    target.sheet_properties.tabColor = copy(source.sheet_properties.tabColor)

                column_map: dict[int, int] = {}
                target_column = 0
                for source_column in range(1, source.max_column + 1):
                    if source_column == excluded_column:
                        continue
                    target_column += 1
                    column_map[source_column] = target_column
                    for row_number in range(1, source.max_row + 1):
                        source_cell = source.cell(row_number, source_column)
                        target_cell = target.cell(row_number, target_column, source_cell.value)
                        if source_cell.has_style:
                            target_cell.font = copy(source_cell.font)
                            target_cell.fill = copy(source_cell.fill)
                            target_cell.border = copy(source_cell.border)
                            target_cell.alignment = copy(source_cell.alignment)
                            target_cell.protection = copy(source_cell.protection)
                            target_cell.number_format = source_cell.number_format
                        if source_cell.hyperlink:
                            target_cell._hyperlink = copy(source_cell.hyperlink)

                for row_number, source_dimension in source.row_dimensions.items():
                    target_dimension = target.row_dimensions[row_number]
                    target_dimension.height = source_dimension.height
                    target_dimension.hidden = source_dimension.hidden
                    target_dimension.outlineLevel = source_dimension.outlineLevel

                from openpyxl.utils import column_index_from_string, get_column_letter, range_boundaries
                from openpyxl.utils.cell import coordinate_from_string
                from openpyxl.worksheet.table import Table

                for source_column, target_column in column_map.items():
                    source_dimension = source.column_dimensions[get_column_letter(source_column)]
                    target_dimension = target.column_dimensions[get_column_letter(target_column)]
                    target_dimension.width = source_dimension.width
                    target_dimension.hidden = source_dimension.hidden
                    target_dimension.bestFit = source_dimension.bestFit
                    target_dimension.outlineLevel = source_dimension.outlineLevel

                for merged_range in source.merged_cells.ranges:
                    min_col, min_row, max_col, max_row = range_boundaries(str(merged_range))
                    if excluded_column and min_col <= excluded_column <= max_col:
                        continue
                    target.merge_cells(
                        start_row=min_row,
                        start_column=column_map[min_col],
                        end_row=max_row,
                        end_column=column_map[max_col],
                    )

                if source.freeze_panes:
                    coordinate = getattr(source.freeze_panes, "coordinate", str(source.freeze_panes))
                    column_letter, row_number = coordinate_from_string(coordinate)
                    source_column = column_index_from_string(column_letter)
                    if source_column in column_map:
                        target.freeze_panes = f"{get_column_letter(column_map[source_column])}{row_number}"

                for source_table in source.tables.values():
                    min_col, min_row, max_col, max_row = range_boundaries(source_table.ref)
                    if excluded_column and min_col <= excluded_column < max_col:
                        continue
                    if excluded_column == max_col:
                        max_col -= 1
                    if min_col not in column_map or max_col not in column_map:
                        continue
                    table = Table(
                        displayName=source_table.displayName,
                        ref=(
                            f"{get_column_letter(column_map[min_col])}{min_row}:"
                            f"{get_column_letter(column_map[max_col])}{max_row}"
                        ),
                    )
                    table.tableStyleInfo = copy(source_table.tableStyleInfo)
                    target.add_table(table)

                if source.auto_filter.ref:
                    min_col, min_row, max_col, max_row = range_boundaries(source.auto_filter.ref)
                    if excluded_column == max_col:
                        max_col -= 1
                    if min_col in column_map and max_col in column_map:
                        target.auto_filter.ref = (
                            f"{get_column_letter(column_map[min_col])}{min_row}:"
                            f"{get_column_letter(column_map[max_col])}{max_row}"
                        )

                if getattr(exported, "calculation", None) is not None:
                    exported.calculation.fullCalcOnLoad = True
                    exported.calculation.forceFullCalc = True
                content = io.BytesIO()
                exported.save(content)
                exported.close()
                filename = f"ATIVIDADES_{datetime.now():%Y%m%d_%H%M}.xlsx"
                return filename, content.getvalue()
            finally:
                workbook.close()

    @staticmethod
    def _validate_service_calc_sheet(sheet) -> None:
        actual = [normalize_header(sheet.cell(1, column).value) for column in range(1, 11)]
        expected = [normalize_header(value) for value in SERVICE_CALC_HEADERS]
        if actual != expected:
            raise RepositoryError(
                f"A aba '{SERVICE_CALC_SHEET}' já existe, mas não possui a estrutura criada pelo aplicativo."
            )

    @staticmethod
    def _prepare_service_calc_sheet(workbook):
        from openpyxl.styles import Font, PatternFill

        if SERVICE_CALC_SHEET in workbook.sheetnames:
            sheet = workbook[SERVICE_CALC_SHEET]
            LocalExcelRepository._validate_service_calc_sheet(sheet)
            return sheet

        sheet = workbook.create_sheet(SERVICE_CALC_SHEET)
        sheet.append(SERVICE_CALC_HEADERS)
        header_fill = PatternFill("solid", fgColor="DD7FD4")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, color="111827")
        widths = (22, 18, 22, 18, 64, 12, 18, 14, 18, 22)
        from openpyxl.utils import get_column_letter

        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "DD7FD4"
        return sheet

    @staticmethod
    def _validate_material_calc_sheet(sheet) -> None:
        actual = [normalize_header(sheet.cell(1, column).value) for column in range(1, 11)]
        expected = [normalize_header(value) for value in MATERIAL_CALC_HEADERS]
        if actual != expected:
            raise RepositoryError(
                f"A aba '{MATERIAL_CALC_SHEET}' já existe, mas não possui a estrutura criada pelo aplicativo."
            )

    @staticmethod
    def _prepare_material_calc_sheet(workbook):
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter

        if MATERIAL_CALC_SHEET in workbook.sheetnames:
            sheet = workbook[MATERIAL_CALC_SHEET]
            LocalExcelRepository._validate_material_calc_sheet(sheet)
            return sheet

        sheet = workbook.create_sheet(MATERIAL_CALC_SHEET)
        sheet.append(MATERIAL_CALC_HEADERS)
        header_fill = PatternFill("solid", fgColor="DD7FD4")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(bold=True, color="111827")
        widths = (22, 18, 22, 18, 64, 12, 18, 14, 18, 22)
        for column, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(column)].width = width
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.tabColor = "7A2F73"
        return sheet

    @staticmethod
    def _standard_labor_cost(activity: dict[str, Any], settings: dict[str, Any]) -> float:
        daily_rate = max(0.0, to_number(settings.get("global_daily_rate")))
        technicians = max(0.0, to_number(activity.get("qtde_tecnicos")))
        days = max(0.0, to_number(activity.get("qtde_dias")))
        return round(daily_rate * technicians * days, 2)

    def get_service_calculation(self, activity_reference: str) -> dict[str, Any]:
        catalog = self.list_services()
        catalog_by_key = {item["key"]: item for item in catalog["items"]}
        materials_error = ""
        try:
            material_catalog = self.list_materials()
        except RepositoryError as exc:
            material_catalog = {"items": [], "source": "MATERIAL.xlsx"}
            materials_error = str(exc)
        materials_by_key = {item["key"]: item for item in material_catalog["items"]}
        settings = self.get_settings()
        with self._lock:
            openpyxl = self._openpyxl()
            try:
                workbook = self._load_workbook(
                    openpyxl,
                    read_only=True,
                    data_only=False,
                    keep_vba=self.file_path.suffix.casefold() == ".xlsm",
                )
            except (OSError, ValueError) as exc:
                raise RepositoryError(f"Não foi possível abrir a base de atividades: {exc}") from exc

            try:
                if self.sheet_name not in workbook.sheetnames:
                    raise RepositoryError(f"Aba '{self.sheet_name}' não encontrada.")
                activity_sheet = workbook[self.sheet_name]
                _headers, field_columns, mode = self._schema(activity_sheet)
                target_row = self._find_activity_row(activity_sheet, field_columns, activity_reference)
                activity_id = clean_text(activity_sheet.cell(target_row, field_columns["id"]).value, 60)
                raw = {"id": activity_id, **self._row_payload(activity_sheet, field_columns, target_row)}
                raw["calculation_mode"] = mode
                raw["custo_tecnico_dia"] = settings["global_daily_rate"]
                activity = normalize_activity(raw)
                record_key = self._record_key(target_row, activity_id)
                activity["record_key"] = record_key

                selected: list[dict[str, Any]] = []
                if SERVICE_CALC_SHEET in workbook.sheetnames:
                    calc_sheet = workbook[SERVICE_CALC_SHEET]
                    self._validate_service_calc_sheet(calc_sheet)
                    for row in calc_sheet.iter_rows(min_row=2, values_only=True):
                        if clean_text(row[0] if len(row) > 0 else "", 200) != record_key:
                            continue
                        service_key = clean_text(row[2] if len(row) > 2 else "", 120)
                        quantity = round(to_number(row[7] if len(row) > 7 else 0), 4)
                        current_service = catalog_by_key.get(service_key)
                        unit_price = round(
                            to_number(
                                current_service["unit_price"]
                                if current_service
                                else row[6] if len(row) > 6 else 0
                            ),
                            4,
                        )
                        selected.append(
                            {
                                "key": service_key,
                                "code": current_service["code"] if current_service else clean_text(row[3] if len(row) > 3 else "", 120),
                                "description": current_service["description"] if current_service else clean_text(row[4] if len(row) > 4 else "", 500),
                                "unit": current_service["unit"] if current_service else clean_text(row[5] if len(row) > 5 else "", 40),
                                "unit_price": unit_price,
                                "quantity": quantity,
                                "subtotal": round(unit_price * quantity, 2),
                            }
                        )

                selected_materials: list[dict[str, Any]] = []
                if MATERIAL_CALC_SHEET in workbook.sheetnames:
                    material_sheet = workbook[MATERIAL_CALC_SHEET]
                    self._validate_material_calc_sheet(material_sheet)
                    for row in material_sheet.iter_rows(min_row=2, values_only=True):
                        if clean_text(row[0] if len(row) > 0 else "", 200) != record_key:
                            continue
                        material_key = clean_text(row[2] if len(row) > 2 else "", 120)
                        quantity = round(to_number(row[7] if len(row) > 7 else 0), 4)
                        current_material = materials_by_key.get(material_key)
                        unit_price = round(
                            to_number(
                                current_material["unit_price"]
                                if current_material
                                else row[6] if len(row) > 6 else 0
                            ),
                            4,
                        )
                        selected_materials.append(
                            {
                                "key": material_key,
                                "code": current_material["code"] if current_material else clean_text(row[3] if len(row) > 3 else "", 120),
                                "description": current_material["description"] if current_material else clean_text(row[4] if len(row) > 4 else "", 500),
                                "unit": current_material["unit"] if current_material else clean_text(row[5] if len(row) > 5 else "", 40),
                                "unit_price": unit_price,
                                "quantity": quantity,
                                "subtotal": round(unit_price * quantity, 2),
                            }
                        )

                service_total = round(sum(to_number(item["subtotal"]) for item in selected), 2)
                material_total = round(sum(to_number(item["subtotal"]) for item in selected_materials), 2)
                standard_labor_cost = self._standard_labor_cost(activity, settings)
                return {
                    "activity": {
                        "id": activity["id"],
                        "record_key": record_key,
                        "tipo_atividade": activity["tipo_atividade"],
                        "custo_mo": activity["custo_mo"],
                        "custo_mat": activity["custo_mat"],
                        "custo_total": activity["custo_total"],
                        "gap": activity["gap"],
                        "qtde_tecnicos": activity["qtde_tecnicos"],
                        "qtde_dias": activity["qtde_dias"],
                    },
                    "services": catalog["items"],
                    "selected": selected,
                    "selected_total": service_total,
                    "materials": material_catalog["items"],
                    "selected_materials": selected_materials,
                    "selected_material_total": material_total,
                    "material_total": material_total if selected_materials else activity["custo_mat"],
                    "material_source": material_catalog["source"],
                    "materials_error": materials_error,
                    "standard_daily_rate": round(to_number(settings.get("global_daily_rate")), 2),
                    "standard_labor_cost": standard_labor_cost,
                    "labor_total": service_total if selected else standard_labor_cost,
                    "gap": round((service_total if selected else standard_labor_cost) - standard_labor_cost, 2),
                    "source": catalog["source"],
                }
            finally:
                workbook.close()

    def save_service_calculation(self, activity_reference: str, payload: dict[str, Any]) -> dict[str, Any]:
        if not isinstance(payload, dict) or not isinstance(payload.get("items"), list):
            raise ValueError("Lista de serviços inválida.")
        if len(payload["items"]) > 500:
            raise ValueError("A calculadora aceita no máximo 500 serviços por atividade.")
        materials_changed = payload.get("materials_changed") is True
        raw_material_items = payload.get("material_items", [])
        if materials_changed and not isinstance(raw_material_items, list):
            raise ValueError("Lista de materiais inválida.")
        if materials_changed and len(raw_material_items) > 500:
            raise ValueError("A calculadora aceita no máximo 500 materiais por atividade.")

        def optional_positive_integer(field: str, label: str) -> int | None:
            if field not in payload or payload.get(field) in (None, ""):
                return None
            raw_value = to_number(payload.get(field))
            if raw_value <= 0 or raw_value > 1_000_000 or not raw_value.is_integer():
                raise ValueError(f"{label} deve ser um número inteiro maior que zero.")
            return int(raw_value)

        technicians = optional_positive_integer("technicians", "A quantidade de técnicos")
        days = optional_positive_integer("days", "A quantidade de dias")

        catalog = self.list_services()
        settings = self.get_settings()
        by_key = {item["key"]: item for item in catalog["items"]}
        quantities: dict[str, int] = {}
        for raw_item in payload["items"]:
            if not isinstance(raw_item, dict):
                raise ValueError("A lista contém um serviço inválido.")
            service_key = clean_text(raw_item.get("key"), 120)
            if service_key not in by_key:
                raise ValueError(
                    "Um serviço selecionado não existe mais na planilha SERVICOS.xlsx. Atualize a calculadora."
                )
            raw_quantity = to_number(raw_item.get("quantity"))
            if raw_quantity <= 0 or raw_quantity > 1_000_000_000 or not raw_quantity.is_integer():
                raise ValueError("A quantidade de cada serviço deve ser um número inteiro maior que zero.")
            quantity = int(raw_quantity)
            quantities[service_key] = int(quantities.get(service_key, 0) + quantity)

        selected: list[dict[str, Any]] = []
        for service_key, quantity in quantities.items():
            service = by_key[service_key]
            subtotal = round(to_number(service["unit_price"]) * quantity, 2)
            selected.append({**service, "quantity": quantity, "subtotal": subtotal})
        service_total = round(sum(item["subtotal"] for item in selected), 2)

        selected_materials: list[dict[str, Any]] = []
        material_catalog = {"items": [], "source": "MATERIAL.xlsx"}
        if materials_changed:
            material_catalog = self.list_materials()
            materials_by_key = {item["key"]: item for item in material_catalog["items"]}
            material_quantities: dict[str, float] = {}
            for raw_item in raw_material_items:
                if not isinstance(raw_item, dict):
                    raise ValueError("A lista contém um material inválido.")
                material_key = clean_text(raw_item.get("key"), 120)
                if material_key not in materials_by_key:
                    raise ValueError(
                        "Um material selecionado não existe mais na planilha MATERIAL.xlsx. Atualize a calculadora."
                    )
                quantity = round(to_number(raw_item.get("quantity")), 4)
                if quantity <= 0 or quantity > 1_000_000_000:
                    raise ValueError("A quantidade de cada material deve ser maior que zero.")
                material_quantities[material_key] = round(
                    material_quantities.get(material_key, 0.0) + quantity,
                    4,
                )

            for material_key, quantity in material_quantities.items():
                material = materials_by_key[material_key]
                subtotal = round(to_number(material["unit_price"]) * quantity, 2)
                selected_materials.append({**material, "quantity": quantity, "subtotal": subtotal})

        with self._lock:
            openpyxl = self._openpyxl()
            workbook = None
            try:
                workbook = self._load_workbook(
                    openpyxl,
                    keep_vba=self.file_path.suffix.casefold() == ".xlsm",
                )
                if self.sheet_name not in workbook.sheetnames:
                    raise RepositoryError(f"Aba '{self.sheet_name}' não encontrada.")
                activity_sheet = workbook[self.sheet_name]
                self._ensure_draft_column(activity_sheet)
                headers, field_columns, mode = self._schema(activity_sheet)
                if "custo_mo" not in field_columns:
                    raise RepositoryError("A coluna Custo MO não foi encontrada na base de atividades.")
                target_row = self._find_activity_row(activity_sheet, field_columns, activity_reference)
                activity_id = clean_text(activity_sheet.cell(target_row, field_columns["id"]).value, 60)
                record_key = self._record_key(target_row, activity_id)

                merged_payload = self._row_payload(activity_sheet, field_columns, target_row)
                if technicians is not None:
                    merged_payload["qtde_tecnicos"] = technicians
                if days is not None:
                    merged_payload["qtde_dias"] = days
                merged_payload["custo_tecnico_dia"] = settings["global_daily_rate"]
                standard_labor_cost = self._standard_labor_cost(merged_payload, settings)
                labor_total = service_total if selected else standard_labor_cost
                merged_payload["custo_mo"] = labor_total
                merged_payload["servico_mo"] = "; ".join(item["description"] for item in selected)
                merged_payload["qtde_servico"] = sum(item["quantity"] for item in selected)
                if materials_changed:
                    material_total = round(sum(item["subtotal"] for item in selected_materials), 2)
                    merged_payload["custo_mat"] = material_total
                    merged_payload["material_utilizado"] = "; ".join(
                        f"{item['description']} ({item['quantity']:g} {item['unit'] or 'un.'})"
                        for item in selected_materials
                    )
                    merged_payload["cod_material"] = "; ".join(
                        item["code"] for item in selected_materials if item["code"]
                    )
                    merged_payload["quantidade"] = round(
                        sum(to_number(item["quantity"]) for item in selected_materials),
                        4,
                    )
                else:
                    material_total = round(to_number(merged_payload.get("custo_mat")), 2)
                activity = self._prepare_activity(activity_id, merged_payload, mode)
                self._apply_activity_to_row(
                    activity_sheet,
                    headers,
                    field_columns,
                    target_row,
                    activity,
                    mode,
                    include_id=False,
                )

                calc_sheet = self._prepare_service_calc_sheet(workbook)
                for row_number in range(calc_sheet.max_row, 1, -1):
                    if clean_text(calc_sheet.cell(row_number, 1).value, 200) == record_key:
                        calc_sheet.delete_rows(row_number, 1)

                updated_at = datetime.now()
                for item in selected:
                    calc_sheet.append(
                        [
                            record_key,
                            int(activity_id) if activity_id.isdigit() else activity_id,
                            item["key"],
                            item["code"],
                            item["description"],
                            item["unit"],
                            item["unit_price"],
                            item["quantity"],
                            item["subtotal"],
                            updated_at,
                        ]
                    )
                    row_number = calc_sheet.max_row
                    calc_sheet.cell(row_number, 7).number_format = 'R$ #,##0.00'
                    calc_sheet.cell(row_number, 8).number_format = '#,##0'
                    calc_sheet.cell(row_number, 9).number_format = 'R$ #,##0.00'
                    calc_sheet.cell(row_number, 10).number_format = 'dd/mm/yyyy hh:mm'
                calc_sheet.auto_filter.ref = f"A1:J{max(1, calc_sheet.max_row)}"

                if materials_changed:
                    material_sheet = self._prepare_material_calc_sheet(workbook)
                    for row_number in range(material_sheet.max_row, 1, -1):
                        if clean_text(material_sheet.cell(row_number, 1).value, 200) == record_key:
                            material_sheet.delete_rows(row_number, 1)
                    for item in selected_materials:
                        material_sheet.append(
                            [
                                record_key,
                                int(activity_id) if activity_id.isdigit() else activity_id,
                                item["key"],
                                item["code"],
                                item["description"],
                                item["unit"],
                                item["unit_price"],
                                item["quantity"],
                                item["subtotal"],
                                updated_at,
                            ]
                        )
                        row_number = material_sheet.max_row
                        material_sheet.cell(row_number, 7).number_format = 'R$ #,##0.00'
                        material_sheet.cell(row_number, 8).number_format = '#,##0.####'
                        material_sheet.cell(row_number, 9).number_format = 'R$ #,##0.00'
                        material_sheet.cell(row_number, 10).number_format = 'dd/mm/yyyy hh:mm'
                    material_sheet.auto_filter.ref = f"A1:J{max(1, material_sheet.max_row)}"

                self._save_workbook(workbook)
                workbook = None
                saved = normalize_activity(activity)
                saved["record_key"] = record_key
                return {
                    "activity": saved,
                    "selected": selected,
                    "selected_total": service_total,
                    "selected_materials": selected_materials,
                    "material_total": material_total,
                    "standard_daily_rate": round(to_number(settings.get("global_daily_rate")), 2),
                    "standard_labor_cost": standard_labor_cost,
                    "labor_total": labor_total,
                    "gap": activity["gap"],
                    "source": catalog["source"],
                    "material_source": material_catalog["source"],
                }
            except PermissionError as exc:
                raise RepositoryError(
                    "O Excel está bloqueado. Feche B2B_CTACUSTOS.xlsx e tente salvar novamente."
                ) from exc
            except OSError as exc:
                raise RepositoryError(f"Não foi possível salvar o cálculo de serviços: {exc}") from exc
            finally:
                if workbook is not None:
                    workbook.close()

    def get_settings(self) -> dict[str, Any]:
        activity_options = self._activity_options(self.list_activities())
        return self._read_settings(activity_options)

    def save_settings(self, payload: dict[str, Any]) -> dict[str, Any]:
        if not isinstance(payload, dict):
            raise ValueError("Configuração inválida.")

        reference_month = clean_text(payload.get("reference_month"), 7) or date.today().strftime("%Y-%m")
        try:
            datetime.strptime(reference_month, "%Y-%m")
        except ValueError as exc:
            raise ValueError("Mês de referência inválido.") from exc
        business_days = int(to_number(payload.get("business_days"))) or business_days_for_month(reference_month)
        if "monthly_technician_cost" in payload:
            monthly_cost = to_number(payload.get("monthly_technician_cost"))
            global_rate = round(monthly_cost / business_days, 2) if business_days > 0 else 0.0
        else:
            # Compatibility with the previous API, which accepted only a daily rate.
            global_rate = to_number(payload.get("global_daily_rate"))
            monthly_cost = round(global_rate * business_days, 2) if global_rate > 0 else 15_000.0
            global_rate = global_rate or round(monthly_cost / business_days, 2)
        if monthly_cost < 0 or business_days <= 0:
            raise ValueError("Valores de configuração inválidos.")

        def clean_list(field: str, max_length: int = 160) -> list[str]:
            values = payload.get(field, [])
            if not isinstance(values, list):
                raise ValueError(f"O cadastro de {field} é inválido.")
            unique: dict[str, str] = {}
            for value in values:
                text = clean_text(value, max_length)
                if text:
                    unique.setdefault(text.casefold(), text)
            return list(unique.values())

        statuses = clean_list("statuses")
        current_settings = self.get_settings()
        types = (
            clean_list("types")
            if "types" in payload
            else [clean_text(value, 160) for value in current_settings.get("types", []) if clean_text(value, 160)]
        )

        raw_technology_rates = payload.get("technology_rates")
        if raw_technology_rates is None and "technologies" in payload:
            raw_technology_rates = [
                {"name": name, "service_cost": 0}
                for name in clean_list("technologies")
            ]
        if raw_technology_rates is None:
            raw_technology_rates = current_settings.get("technology_rates", [])
        if not isinstance(raw_technology_rates, list):
            raise ValueError("O cadastro de tecnologias é inválido.")
        technology_map: dict[str, dict[str, Any]] = {}
        for item in raw_technology_rates:
            if isinstance(item, str):
                name = clean_text(item, 160)
                service_cost = 0.0
            elif isinstance(item, dict):
                name = clean_text(item.get("name"), 160)
                service_cost = to_number(item.get("service_cost"))
            else:
                raise ValueError("O cadastro de tecnologias contém um item inválido.")
            if service_cost < 0:
                raise ValueError("O valor do serviço por tecnologia não pode ser negativo.")
            if name:
                technology_map.setdefault(
                    name.casefold(),
                    {"name": name, "service_cost": round(service_cost, 2)},
                )
        technology_rates = list(technology_map.values())
        technologies = [item["name"] for item in technology_rates]

        raw_category_teams = payload.get("category_teams", current_settings.get("category_teams", []))
        if not isinstance(raw_category_teams, list):
            raise ValueError("A configuração das equipes por categoria é inválida.")
        supplied_teams: dict[str, int] = {}
        valid_category_keys = {str(item["key"]) for item in SERVICE_CATEGORIES}
        for item in raw_category_teams:
            if not isinstance(item, dict):
                raise ValueError("A configuração das equipes contém um item inválido.")
            category_key = service_category(item.get("category")) or normalize_header(item.get("category"))
            raw_technicians = to_number(item.get("technicians"))
            if category_key not in valid_category_keys or raw_technicians < 0 or not raw_technicians.is_integer():
                raise ValueError("Informe quantidades inteiras e não negativas para as equipes por categoria.")
            supplied_teams[category_key] = int(raw_technicians)
        category_teams = [
            {
                "category": category["key"],
                "label": category["label"],
                "technicians": supplied_teams.get(str(category["key"]), 0),
            }
            for category in SERVICE_CATEGORIES
        ]
        companies = clean_list("companies")
        eps_values = clean_list("eps")
        raw_technicians = payload.get("technicians", [])
        if not isinstance(raw_technicians, list):
            raise ValueError("O cadastro de técnicos é inválido.")
        technician_map: dict[str, dict[str, str]] = {}
        for item in raw_technicians:
            if not isinstance(item, dict):
                raise ValueError("O cadastro de técnicos contém um item inválido.")
            registration = clean_text(item.get("registration"), 120)
            name = clean_text(item.get("name"), 160)
            if not registration and not name:
                continue
            key = (f"registration:{registration}" if registration else f"name:{name}").casefold()
            technician_map.setdefault(key, {"registration": registration, "name": name})
        technicians = list(technician_map.values())

        raw_uploads = payload.get("uploads")
        if not isinstance(raw_uploads, dict):
            raw_uploads = self.get_settings().get("uploads", {})
        uploads: dict[str, dict[str, str]] = {}
        for kind in ("services", "materials"):
            item = raw_uploads.get(kind, {}) if isinstance(raw_uploads, dict) else {}
            if not isinstance(item, dict):
                item = {}
            uploads[kind] = {
                "filename": clean_text(item.get("filename"), 240),
                "path": clean_text(item.get("path"), 500),
                "uploaded_at": clean_text(item.get("uploaded_at"), 40),
            }

        with self._lock:
            openpyxl = self._openpyxl()
            try:
                workbook = self._load_workbook(openpyxl, keep_vba=self.file_path.suffix.casefold() == ".xlsm")
                sheet = workbook["Config"] if "Config" in workbook.sheetnames else workbook.create_sheet("Config")
                if sheet.max_row:
                    sheet.delete_rows(1, sheet.max_row)
                sheet.append(["Configuração de custos"])
                sheet.append(["Custo mensal técnico", monthly_cost])
                sheet.append(["Dias úteis do mês", business_days])
                sheet.append(["Custo técnico/dia", global_rate])
                sheet.append(["Mês de referência", reference_month])
                sheet.append(["Pasta de importações", str(self._reference_upload_dir())])
                sheet.append(["Uso", "Indicadores, cadastros e planilhas de referência do B2B CTACUSTOS."])
                sheet.append(["Categoria", "Código/Matrícula", "Nome/Valor", "Atualizado em"])
                for status in statuses:
                    sheet.append(["STATUS", "", status])
                for activity_type in types:
                    sheet.append(["TIPO_ATIVIDADE", "", activity_type])
                if not types:
                    sheet.append(["TIPO_ATIVIDADE", "", ""])
                for technology in technology_rates:
                    sheet.append(["TECNOLOGIA", technology["service_cost"], technology["name"]])
                if not technology_rates:
                    sheet.append(["TECNOLOGIA", "", ""])
                for team in category_teams:
                    sheet.append(["EQUIPE_CATEGORIA", team["category"], team["technicians"]])
                for item in technicians:
                    sheet.append(["TECNICO", item["registration"], item["name"]])
                for company in companies:
                    sheet.append(["EMPRESA", "", company])
                for eps in eps_values:
                    sheet.append(["EPS", "", eps])
                for kind, category in (("services", "ARQUIVO_SERVICOS"), ("materials", "ARQUIVO_MATERIAIS")):
                    item = uploads[kind]
                    if item["filename"] or item["path"]:
                        sheet.append([category, item["filename"], item["path"], item["uploaded_at"]])

                from openpyxl.styles import Font, PatternFill

                header_fill = PatternFill("solid", fgColor="DD7FD4")
                for cell in (*sheet[1], *sheet[8]):
                    cell.fill = header_fill
                    cell.font = Font(bold=True, color="111827")
                sheet["B2"].number_format = 'R$ #,##0.00'
                sheet["B3"].number_format = "0"
                sheet["B4"].number_format = 'R$ #,##0.00'
                for row_number in range(9, sheet.max_row + 1):
                    if normalize_header(sheet.cell(row_number, 1).value) == "tecnologia":
                        sheet.cell(row_number, 2).number_format = 'R$ #,##0.00'
                    elif normalize_header(sheet.cell(row_number, 1).value) == "equipecategoria":
                        sheet.cell(row_number, 3).number_format = '0'
                sheet.column_dimensions["A"].width = 24
                sheet.column_dimensions["B"].width = 24
                sheet.column_dimensions["C"].width = 58
                sheet.column_dimensions["D"].width = 22
                sheet.freeze_panes = "A9"
                sheet.sheet_view.showGridLines = False
                sheet.auto_filter.ref = f"A8:D{max(8, sheet.max_row)}"
                self._save_workbook(workbook)
                return self.get_settings()
            except PermissionError as exc:
                raise RepositoryError("O Excel está bloqueado para salvar configurações. Feche o arquivo e tente novamente.") from exc
            except OSError as exc:
                raise RepositoryError(f"Não foi possível salvar as configurações: {exc}") from exc

    def save_reference_workbook(self, kind: str, filename: str, content: bytes) -> dict[str, str]:
        labels = {"services": "Serviços", "materials": "Materiais"}
        if kind not in labels:
            raise ValueError("Tipo de planilha inválido.")
        original_name = Path(str(filename or "").replace("\\", "/")).name
        safe_name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", original_name).strip(" .")
        suffix = Path(safe_name).suffix.casefold()
        if not safe_name or suffix not in {".xlsx", ".xlsm"}:
            raise ValueError("Selecione uma planilha .xlsx ou .xlsm.")
        if not content or len(content) > 25 * 1024 * 1024:
            raise ValueError("A planilha deve ter no máximo 25 MB.")
        if not zipfile.is_zipfile(io.BytesIO(content)):
            raise ValueError("O arquivo enviado não é uma planilha Excel válida.")

        directory = self._reference_upload_dir()
        with self._lock:
            try:
                directory.mkdir(parents=True, exist_ok=True)
                destination = directory / safe_name
                if destination.exists():
                    stamp = datetime.now().strftime("%Y%m%d-%H%M%S-%f")
                    destination = directory / f"{destination.stem}-{stamp}{destination.suffix}"
                temporary = directory / f".{destination.name}.{uuid.uuid4().hex}.uploading"
                try:
                    temporary.write_bytes(content)
                    os.replace(temporary, destination)
                finally:
                    if temporary.exists():
                        temporary.unlink()

                metadata = {
                    "filename": destination.name,
                    "path": str(destination),
                    "uploaded_at": datetime.now().isoformat(timespec="seconds"),
                }
                settings = self.get_settings()
                settings["uploads"][kind] = metadata
                try:
                    self.save_settings(settings)
                except Exception:
                    try:
                        destination.unlink()
                    except OSError:
                        pass
                    raise
                return {**metadata, "directory": str(directory), "label": labels[kind]}
            except PermissionError as exc:
                raise RepositoryError(
                    "Não foi possível gravar na pasta sincronizada. Verifique a permissão do OneDrive."
                ) from exc
            except OSError as exc:
                raise RepositoryError(f"Não foi possível salvar a planilha de {labels[kind].lower()}: {exc}") from exc


def build_repository(base_dir: Path):
    return LocalExcelRepository(base_dir)
