import io
import os
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from unittest.mock import patch

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))

from excel_repository import (  # noqa: E402
    LocalExcelRepository,
    RepositoryError,
    business_days_for_month,
    calculate_values,
    discover_excel_file,
    normalize_activity,
    summarize_activities,
    summarize_service_categories,
)
from app import load_env  # noqa: E402


def create_test_activity_workbook(path: Path) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    from openpyxl.worksheet.table import Table, TableStyleInfo

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Atividades"
    sheet.append(
        [
            "ID", "Tipo de Atividade", "Status", "Situação", "Material Utilizado",
            "Cod. Material", "Quantidade", "Custo Mat", "Serviço M.O", "Qtde Serviço",
            "Custo M.O", "Custo Total", "Custo Evitado", "Custo Técnico / Dia",
            "Qtde Técnicos", "Qtde Dias", "Custo por Técnico", "GAP", "Atualizado Em",
        ]
    )
    statuses = ["Planejada", "Em andamento", "Em validação", "Concluída", "Cancelada"]
    for index, status in enumerate(statuses, start=1):
        sheet.append(
            [
                f"ATV-{index:03d}", f"Atividade de teste {index}", status,
                f"Situação de teste {index}", "Material de teste", f"MAT-{index:03d}",
                index, 10 * index, "Serviço de teste", index, 20 * index,
                30 * index, 50 * index, 100, 1, 1, 30 * index, 20 * index,
                f"2026-09-{index:02d}",
            ]
        )
    for cell in sheet[1]:
        cell.fill = PatternFill("solid", fgColor="DD7FD4")
        cell.font = Font(bold=True, color="18151A")
    table = Table(displayName="AtividadesTestTable", ref=f"A1:S{sheet.max_row}")
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium4",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    sheet.add_table(table)
    workbook.save(path)
    workbook.close()


class CalculationTests(unittest.TestCase):
    def test_local_env_replaces_stale_process_value(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            env_path = Path(temp_dir) / ".env"
            env_path.write_text("ADMIN_PASSWORD=new-value\n", encoding="utf-8")
            with patch.dict(os.environ, {"ADMIN_PASSWORD": "old-value"}, clear=False):
                load_env(env_path)
                self.assertEqual(os.environ["ADMIN_PASSWORD"], "new-value")

    def test_activity_modal_has_only_requested_cost_fields(self):
        html = (PROJECT / "web" / "atividades.html").read_text(encoding="utf-8")
        javascript = (PROJECT / "web" / "static" / "app.js").read_text(encoding="utf-8")
        for field in ("custo_mo", "custo_mat", "custo_evitado"):
            self.assertIn(f'name="{field}"', html)
        self.assertIn('id="calc-total"', html)
        self.assertIn('id="form-service-cost"', html)
        self.assertIn('normalize(technology?.value).trim() === "erb"', javascript)
        for field in ("quantidade", "servico_mo", "custo_tecnico_dia"):
            self.assertNotIn(f'name="{field}"', html)
        self.assertIn('id="calc-gap"', html)
        self.assertIn('id="service-gap"', html)
        self.assertIn('id="service-calculator-modal"', html)
        self.assertIn('id="service-calculator-form"', html)
        self.assertIn('class="calculator-column"', html)
        self.assertIn('id="calc-avoided"', html)
        self.assertIn('data-currency-input', html)
        self.assertIn('id="service-technicians"', html)
        self.assertIn('id="service-days"', html)
        self.assertIn('name="draft"', html)
        self.assertIn('id="export-button"', html)
        self.assertIn('id="form-technology"', html)
        self.assertIn('Custo Serviço', html)
        self.assertNotIn('name="ganho_esperado"', html)
        self.assertNotIn('valor técnico/dia × técnicos × dias neste formulário', html)

    def test_category_summary_uses_configured_monthly_team_values(self):
        rows = [
            {"tipo_atividade": "IMPLANTAÇÃO", "custo_mo": 4_000, "custo_mat": 500},
            {"tipo_atividade": "REPARO B2B", "custo_mo": 2_500, "custo_mat": 250},
            {"tipo_atividade": "ATIVAÇÃO B2B", "custo_mo": 7_000, "custo_mat": 700},
        ]
        settings = {
            "monthly_technician_cost": 15_000,
            "category_teams": [
                {"category": "implantacao", "technicians": 2},
                {"category": "reparo", "technicians": 1},
                {"category": "ativacao", "technicians": 3},
            ],
        }
        categories = {item["key"]: item for item in summarize_service_categories(rows, settings)}
        self.assertEqual(categories["implantacao"]["team_value"], 30_000)
        self.assertEqual(categories["implantacao"]["gap"], -26_000)
        self.assertEqual(categories["reparo"]["total"], 2_750)
        self.assertEqual(categories["ativacao"]["technicians"], 3)

    def test_service_catalog_accepts_blank_codes(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            activity_path = temp_root / "B2B_CTACUSTOS.xlsx"
            services_path = temp_root / "SERVICOS.xlsx"

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            sheet.append(["ID", "Obra Executada", "Status Obra", "Detalhes Obra", "Custo MO", "Custo Material", "Custo Total"])
            sheet.append([10, "REPARO", "ABERTO", "EM CAMPO", 0, 0, 0])
            workbook.save(activity_path)
            workbook.close()

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "SERVICOS"
            sheet.append(["CODIGO", "SERVICO", "CUSTO_UNITARIO", "UNIDADE"])
            sheet.append([None, "Instalação de cabo óptico", 1.8, "M"])
            sheet.append([None, "Emenda de FO", 20.32, "UN"])
            workbook.save(services_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_PATH": str(activity_path),
                    "EXCEL_SHEET_NAME": "Atividades",
                    "SERVICES_EXCEL_PATH": str(services_path),
                    "SERVICES_EXCEL_SHEET_NAME": "SERVICOS",
                },
                clear=False,
            ):
                catalog = LocalExcelRepository(temp_root).list_services()

            self.assertEqual(catalog["source"], "SERVICOS.xlsx · aba SERVICOS")
            self.assertEqual(len(catalog["items"]), 2)
            self.assertEqual(catalog["items"][0]["code"], "")
            self.assertEqual(catalog["items"][0]["unit_price"], 1.8)
            self.assertNotEqual(catalog["items"][0]["key"], catalog["items"][1]["key"])

    def test_service_calculation_updates_exact_duplicate_id_row(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            activity_path = temp_root / "B2B_CTACUSTOS.xlsx"
            services_path = temp_root / "SERVICOS.xlsx"

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            sheet.append(
                [
                    "ID", "Obra Executada", "Status Obra", "Detalhes Obra",
                    "Custo MO", "Custo Material", "Custo Total", "Custo Evitado",
                ]
            )
            sheet.append([665232, "REPARO A", "ABERTO", "PRIMEIRA LINHA", 100, 50, 150, 300])
            sheet.append([665232, "REPARO B", "ABERTO", "SEGUNDA LINHA", 200, 40, 240, 400])
            workbook.save(activity_path)
            workbook.close()

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "SERVICOS"
            sheet.append(["CODIGO", "SERVICO", "CUSTO_UNITARIO", "UNIDADE"])
            sheet.append([None, "Instalação de cabo óptico", 1.8, "M"])
            sheet.append([None, "Emenda de FO", 20.32, "UN"])
            workbook.save(services_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_PATH": str(activity_path),
                    "EXCEL_SHEET_NAME": "Atividades",
                    "SERVICES_EXCEL_PATH": str(services_path),
                    "SERVICES_EXCEL_SHEET_NAME": "SERVICOS",
                },
                clear=False,
            ):
                repository = LocalExcelRepository(temp_root)
                activities = repository.list_activities()
                second = next(item for item in activities if item["situacao"] == "SEGUNDA LINHA")
                catalog = repository.list_services()["items"]
                result = repository.save_service_calculation(
                    second["record_key"],
                    {
                        "items": [
                            {"key": catalog[0]["key"], "quantity": 2},
                            {"key": catalog[1]["key"], "quantity": 3},
                        ]
                    },
                )
                calculation = repository.get_service_calculation(second["record_key"])

            expected_labor = round(1.8 * 2 + 20.32 * 3, 2)
            self.assertEqual(result["selected_total"], expected_labor)
            self.assertEqual(calculation["selected_total"], expected_labor)
            self.assertEqual(len(calculation["selected"]), 2)

            workbook = openpyxl.load_workbook(activity_path, data_only=False)
            activities_sheet = workbook["Atividades"]
            self.assertEqual(activities_sheet["E2"].value, 100)
            self.assertEqual(activities_sheet["G2"].value, 150)
            self.assertEqual(activities_sheet["E3"].value, expected_labor)
            self.assertEqual(activities_sheet["G3"].value, expected_labor + 40)
            self.assertEqual(activities_sheet["H3"].value, expected_labor)
            calculation_sheet = workbook["Calculo_Servicos"]
            self.assertEqual(calculation_sheet.max_row, 3)
            self.assertTrue(all(calculation_sheet.cell(row, 1).value == second["record_key"] for row in (2, 3)))
            workbook.close()

    def test_empty_service_calculation_uses_standard_team_cost(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            activity_path = temp_root / "B2B_CTACUSTOS.xlsx"
            services_path = temp_root / "SERVICOS.xlsx"

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            sheet.append(
                [
                    "ID", "Obra Executada", "Status Obra", "Detalhes Obra",
                    "Custo MO", "Custo Material", "Custo Total", "Custo Evitado",
                    "Vol. Técnicos", "Tempo Execução",
                ]
            )
            sheet.append([1, "REPARO", "ABERTO", "EM CAMPO", 100, 50, 150, 300, 1, 1])
            workbook.save(activity_path)
            workbook.close()

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "SERVICOS"
            sheet.append(["CODIGO", "SERVICO", "CUSTO_UNITARIO", "UNIDADE"])
            sheet.append([None, "Emenda de FO", 20.32, "UN"])
            workbook.save(services_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_PATH": str(activity_path),
                    "EXCEL_SHEET_NAME": "Atividades",
                    "SERVICES_EXCEL_PATH": str(services_path),
                    "SERVICES_EXCEL_SHEET_NAME": "SERVICOS",
                },
                clear=False,
            ):
                repository = LocalExcelRepository(temp_root)
                repository.save_settings(
                    {
                        "reference_month": "2026-09",
                        "monthly_technician_cost": 15_000,
                        "business_days": 20,
                        "statuses": [],
                        "technicians": [],
                        "companies": [],
                        "eps": [],
                    }
                )
                activity = repository.list_activities()[0]
                self.assertEqual(activity["gap"], -650)
                service = repository.list_services()["items"][0]
                with self.assertRaisesRegex(ValueError, "número inteiro"):
                    repository.save_service_calculation(
                        activity["record_key"],
                        {"items": [{"key": service["key"], "quantity": 1.5}]},
                    )
                with self.assertRaisesRegex(ValueError, "quantidade de técnicos"):
                    repository.save_service_calculation(
                        activity["record_key"],
                        {"technicians": 1.5, "days": 3, "items": []},
                    )

                result = repository.save_service_calculation(
                    activity["record_key"],
                    {"technicians": 2, "days": 3, "items": []},
                )
                calculation = repository.get_service_calculation(activity["record_key"])

            self.assertEqual(result["selected_total"], 0)
            self.assertEqual(result["standard_daily_rate"], 750)
            self.assertEqual(result["standard_labor_cost"], 4_500)
            self.assertEqual(result["labor_total"], 4_500)
            self.assertEqual(result["gap"], 0)
            self.assertEqual(calculation["selected"], [])
            self.assertEqual(calculation["labor_total"], 4_500)
            self.assertEqual(calculation["gap"], 0)

            workbook = openpyxl.load_workbook(activity_path, data_only=False)
            sheet = workbook["Atividades"]
            self.assertEqual(sheet["E2"].value, 4_500)
            self.assertEqual(sheet["G2"].value, 4_550)
            self.assertEqual(sheet["H2"].value, 4_500)
            self.assertEqual(sheet["I2"].value, 2)
            self.assertEqual(sheet["J2"].value, 3)
            self.assertEqual(workbook["Calculo_Servicos"].max_row, 1)
            workbook.close()

    def test_material_calculation_updates_activity_and_keeps_detailed_memory(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            activity_path = temp_root / "B2B_CTACUSTOS.xlsx"
            services_path = temp_root / "SERVICOS.xlsx"
            materials_path = temp_root / "MATERIAL.xlsx"

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            sheet.append(
                [
                    "ID", "Tipo de Atividade", "Status", "Situação", "Material Utilizado",
                    "Cod. Material", "Quantidade", "Custo Mat", "Serviço M.O", "Qtde Serviço",
                    "Custo M.O", "Custo Total", "Custo Evitado", "Custo Técnico / Dia",
                    "Qtde Técnicos", "Qtde Dias", "Custo por Técnico", "GAP",
                ]
            )
            sheet.append([77, "INSTALAÇÃO", "ABERTO", "EM CAMPO", "", "", 0, 25, "", 0, 100, 125, 100, 500, 1, 1, 125, -400])
            workbook.save(activity_path)
            workbook.close()

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "SERVICOS"
            sheet.append(["CODIGO", "SERVICO", "CUSTO_UNITARIO", "UNIDADE"])
            sheet.append(["S-1", "Serviço padrão", 50, "UN"])
            workbook.save(services_path)
            workbook.close()

            workbook = openpyxl.Workbook()
            sheet = workbook.active
            # The operational file currently uses this legacy sheet name.
            sheet.title = "SERVICOS"
            sheet.append(["Mat_Code", "Material", "Unidade", "Valor"])
            sheet.append(["M-1", "Cabo óptico", "M", 4.5])
            sheet.append(["M-2", "Conector", "UN", 12])
            workbook.save(materials_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_PATH": str(activity_path),
                    "EXCEL_SHEET_NAME": "Atividades",
                    "SERVICES_EXCEL_PATH": str(services_path),
                    "SERVICES_EXCEL_SHEET_NAME": "SERVICOS",
                    "MATERIALS_EXCEL_PATH": str(materials_path),
                    "MATERIALS_EXCEL_SHEET_NAME": "",
                },
                clear=False,
            ):
                repository = LocalExcelRepository(temp_root)
                material_catalog = repository.list_materials()
                activity = repository.list_activities()[0]
                result = repository.save_service_calculation(
                    activity["record_key"],
                    {
                        "technicians": 1,
                        "days": 1,
                        "items": [],
                        "materials_changed": True,
                        "material_items": [
                            {"key": material_catalog["items"][0]["key"], "quantity": 10.5},
                            {"key": material_catalog["items"][1]["key"], "quantity": 2},
                        ],
                    },
                )
                calculation = repository.get_service_calculation(activity["record_key"])

            self.assertEqual(material_catalog["source"], "MATERIAL.xlsx · aba SERVICOS")
            self.assertEqual(result["material_total"], 71.25)
            self.assertEqual(calculation["material_total"], 71.25)
            self.assertEqual(len(calculation["selected_materials"]), 2)

            workbook = openpyxl.load_workbook(activity_path, data_only=False)
            sheet = workbook["Atividades"]
            self.assertEqual(sheet["H2"].value, 71.25)
            self.assertIn("Cabo óptico", sheet["E2"].value)
            self.assertEqual(sheet["F2"].value, "M-1; M-2")
            self.assertEqual(sheet["G2"].value, 12.5)
            self.assertEqual(sheet["L2"].value, "=H2+K2")
            self.assertEqual(workbook["Calculo_Materiais"].max_row, 3)
            workbook.close()

    def test_financial_formula(self):
        result = calculate_values(
            {
                "quantidade": 2,
                "custo_mat": 100,
                "custo_mo": 50,
                "custo_tecnico_dia": 200,
                "qtde_tecnicos": 2,
                "qtde_dias": 2,
                "custo_evitado": 1500,
            }
        )
        self.assertEqual(result["custo_total"], 150)
        self.assertEqual(result["custo_por_tecnico"], 75)
        self.assertEqual(result["gap"], -750)

        positive = calculate_values(
            {
                "custo_mat": 0,
                "custo_mo": 1260,
                "custo_tecnico_dia": 500,
                "qtde_tecnicos": 2,
                "qtde_dias": 1,
            }
        )
        self.assertEqual(positive["gap"], 260)

    def test_business_days_count_uses_monday_to_friday(self):
        self.assertEqual(business_days_for_month("2026-09"), 22)

    def test_summary_uses_calculated_values(self):
        rows = [
            {"status": "Concluída", "tipo_atividade": "A", "tecnologia": "GPON", "custo_mo": 70, "custo_total": 100, "custo_evitado": 180, "gap": 80},
            {"status": "Em andamento", "tipo_atividade": "B", "tecnologia": "B2B", "custo_mo": 60, "custo_total": 90, "custo_evitado": 120, "gap": 30},
        ]
        summary = summarize_activities(rows)
        self.assertEqual(summary["metrics"]["total"], 2)
        self.assertEqual(summary["metrics"]["concluidas"], 1)
        self.assertEqual(summary["financial"]["gap"], 110)
        self.assertEqual(summary["charts"]["by_type"][0]["labor"], 70)
        self.assertEqual({item["label"] for item in summary["charts"]["by_technology"]}, {"GPON", "B2B"})

    def test_sample_workbook_can_be_read_and_updated(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            copy = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(copy)
            old_roots = os.environ.get("EXCEL_SEARCH_ROOTS")
            old_file = os.environ.get("EXCEL_PATH")
            old_filename = os.environ.get("EXCEL_FILENAME")
            try:
                os.environ["EXCEL_SEARCH_ROOTS"] = "data"
                os.environ.pop("EXCEL_PATH", None)
                os.environ["EXCEL_FILENAME"] = "B2B_CTACUSTOS.xlsx"
                repository = LocalExcelRepository(temp_root)
                items = repository.list_activities()
                self.assertGreaterEqual(len(items), 5)
                first = dict(items[0])
                first["situacao"] = "Atualizada pelo teste"
                saved = repository.update_activity(first["id"], first)
                self.assertEqual(saved["situacao"], "Atualizada pelo teste")
                reloaded = repository.list_activities()
                self.assertEqual(reloaded[0]["situacao"], "Atualizada pelo teste")
            finally:
                if old_roots is None:
                    os.environ.pop("EXCEL_SEARCH_ROOTS", None)
                else:
                    os.environ["EXCEL_SEARCH_ROOTS"] = old_roots
                if old_file is None:
                    os.environ.pop("EXCEL_PATH", None)
                else:
                    os.environ["EXCEL_PATH"] = old_file
                if old_filename is None:
                    os.environ.pop("EXCEL_FILENAME", None)
                else:
                    os.environ["EXCEL_FILENAME"] = old_filename

    def test_discovery_refuses_duplicate_workbooks(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            first = temp_root / "OneDrive - A"
            second = temp_root / "OneDrive - B"
            first.mkdir()
            second.mkdir()
            (first / "B2B_CTACUSTOS.xlsx").write_bytes(b"one")
            (second / "B2B_CTACUSTOS.xlsx").write_bytes(b"two")
            old_roots = os.environ.get("EXCEL_SEARCH_ROOTS")
            old_fallback = os.environ.get("EXCEL_FALLBACK_SAMPLE")
            try:
                os.environ["EXCEL_SEARCH_ROOTS"] = os.pathsep.join((str(first), str(second)))
                os.environ["EXCEL_FALLBACK_SAMPLE"] = "false"
                with self.assertRaisesRegex(RepositoryError, "2 cópias"):
                    discover_excel_file("B2B_CTACUSTOS.xlsx", base_dir=temp_root)
            finally:
                if old_roots is None:
                    os.environ.pop("EXCEL_SEARCH_ROOTS", None)
                else:
                    os.environ["EXCEL_SEARCH_ROOTS"] = old_roots
                if old_fallback is None:
                    os.environ.pop("EXCEL_FALLBACK_SAMPLE", None)
                else:
                    os.environ["EXCEL_FALLBACK_SAMPLE"] = old_fallback

    def test_discovery_accepts_plural_reference_alias(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            onedrive = temp_root / "OneDrive"
            onedrive.mkdir()
            expected = onedrive / "MATERIAIS.xlsx"
            expected.write_bytes(b"xlsx-placeholder")
            with patch.dict(
                os.environ,
                {"EXCEL_SEARCH_ROOTS": str(onedrive), "EXCEL_FALLBACK_SAMPLE": "false"},
                clear=False,
            ):
                located = discover_excel_file(
                    "MATERIAL.xlsx",
                    base_dir=temp_root,
                    fallback_filenames=("MATERIAIS.xlsx",),
                )
            self.assertEqual(located, expected.resolve())

    def test_reference_search_ignores_stale_path_and_checks_all_synced_roots(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            restricted_root = temp_root / "Base principal"
            synced_root = temp_root / "OneDrive - Empresa"
            restricted_root.mkdir()
            (synced_root / "Equipe" / "Referencias").mkdir(parents=True)
            main_path = restricted_root / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(main_path)
            reference_name = "SERVICOS_TESTE_PORTABILIDADE.xlsx"
            services_path = synced_root / "Equipe" / "Referencias" / reference_name
            workbook = openpyxl.Workbook()
            workbook.active.title = "SERVIÇOS"
            workbook.save(services_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": str(restricted_root),
                    "OneDriveCommercial": str(synced_root),
                    "EXCEL_FILENAME": main_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                    "SERVICES_EXCEL_PATH": str(temp_root / "caminho-antigo" / reference_name),
                    "SERVICES_EXCEL_FILENAME": reference_name,
                    "SERVICES_EXCEL_FILENAME_ALIASES": "",
                },
                clear=False,
            ):
                repository = LocalExcelRepository(temp_root)
                located = repository._locate_reference_workbook("services")

            self.assertEqual(located, services_path.resolve())

    def test_reference_search_uses_uploaded_filename_from_another_computer(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            main_path = temp_root / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(main_path)
            imports = temp_root / "B2B_CTACUSTOS_Importacoes"
            imports.mkdir()
            expected = imports / "SERVICOS-20260918-101500.xlsx"
            expected.write_bytes(b"xlsx-placeholder")

            with patch.dict(
                os.environ,
                {
                    "EXCEL_PATH": str(main_path),
                    "SERVICES_EXCEL_PATH": "Z:\\outro-computador\\SERVICOS-20260918-101500.xlsx",
                    "SERVICES_EXCEL_FILENAME": "SERVICOS.xlsx",
                },
                clear=False,
            ):
                repository = LocalExcelRepository(temp_root)
                with patch.object(
                    repository,
                    "get_settings",
                    return_value={
                        "uploads": {
                            "services": {
                                "filename": expected.name,
                                "path": "Z:\\outro-computador\\SERVICOS-20260918-101500.xlsx",
                            }
                        }
                    },
                ):
                    located = repository._locate_reference_workbook("services")

            self.assertEqual(located, expected.resolve())

    def test_missing_services_does_not_block_material_catalog(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            main_path = temp_root / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(main_path)

            with patch.dict(
                os.environ,
                {"EXCEL_PATH": str(main_path), "EXCEL_SHEET_NAME": "Atividades"},
                clear=False,
            ):
                repository = LocalExcelRepository(temp_root)
                activity = repository.list_activities()[0]
                with patch.object(repository, "list_services", side_effect=RepositoryError("Serviços indisponíveis")):
                    with patch.object(
                        repository,
                        "list_materials",
                        return_value={
                            "items": [{"key": "mat-1", "code": "1", "description": "Cabo", "unit": "M", "unit_price": 2}],
                            "source": "MATERIAL.xlsx · aba MATERIAIS",
                        },
                    ):
                        calculation = repository.get_service_calculation(activity["record_key"])

            self.assertEqual(calculation["services"], [])
            self.assertEqual(calculation["services_error"], "Serviços indisponíveis")
            self.assertEqual(len(calculation["materials"]), 1)

    def test_record_key_works_when_excel_has_no_stored_sheet_dimension(self):
        class Cell:
            value = "ATV-001"

        class DimensionlessSheet:
            max_row = None

            @staticmethod
            def cell(_row, _column):
                return Cell()

        repository = object.__new__(LocalExcelRepository)
        row = repository._find_activity_row(DimensionlessSheet(), {"id": 1}, "2:ATV-001")
        self.assertEqual(row, 2)

    def test_reference_sheet_name_is_accent_insensitive_and_accepts_only_sheet(self):
        import openpyxl

        workbook = openpyxl.Workbook()
        workbook.active.title = "Serviços"
        with patch.dict(os.environ, {"SERVICES_EXCEL_SHEET_NAME": "SERVICOS"}, clear=False):
            self.assertEqual(LocalExcelRepository._reference_sheet(workbook, "services").title, "Serviços")
        workbook.active.title = "Tabela compartilhada"
        with patch.dict(os.environ, {"SERVICES_EXCEL_SHEET_NAME": "SERVICOS"}, clear=False):
            self.assertEqual(
                LocalExcelRepository._reference_sheet(workbook, "services").title,
                "Tabela compartilhada",
            )
        workbook.close()

    def test_excel_dates_are_normalized_for_date_inputs(self):
        for value in (datetime(2026, 9, 1), "01/09/2026", "2026-09-01T00:00"):
            with self.subTest(value=value):
                self.assertEqual(normalize_activity({"data": value})["data"], "2026-09-01")

    def test_draft_is_saved_and_export_omits_it_without_losing_table_style(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(workbook_path)

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                activity = repository.list_activities()[0]
                activity["draft"] = "DRAFT-001"
                repository.update_activity(activity["record_key"], activity)
                export_name, export_bytes = repository.export_activities()

            source_workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            source_sheet = source_workbook["Atividades"]
            self.assertEqual(source_sheet.cell(1, source_sheet.max_column).value, "DRAFT")
            self.assertEqual(source_sheet.cell(2, source_sheet.max_column).value, "DRAFT-001")
            self.assertEqual(source_sheet.tables["AtividadesTestTable"].ref, "A1:T6")
            self.assertEqual(len(source_sheet.tables["AtividadesTestTable"].tableColumns), 20)
            self.assertEqual(source_sheet.tables["AtividadesTestTable"].tableColumns[-1].name, "DRAFT")
            self.assertEqual(source_sheet.tables["AtividadesTestTable"].autoFilter.ref, "A1:T6")
            source_header_color = source_sheet["A1"].fill.fgColor.rgb
            source_workbook.close()

            exported_workbook = openpyxl.load_workbook(io.BytesIO(export_bytes), data_only=False)
            exported_sheet = exported_workbook["Atividades"]
            exported_headers = [cell.value for cell in exported_sheet[1]]
            self.assertNotIn("DRAFT", exported_headers)
            self.assertEqual(exported_sheet["A1"].fill.fgColor.rgb, source_header_color)
            self.assertEqual(exported_sheet.tables["AtividadesTestTable"].ref, "A1:S6")
            self.assertEqual(len(exported_sheet.tables["AtividadesTestTable"].tableColumns), 19)
            self.assertTrue(export_name.startswith("ATIVIDADES_"))
            exported_workbook.close()

    def test_discovery_uses_current_users_onedrive_environment(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            onedrive = temp_root / "OneDrive - Empresa"
            nested = onedrive / "Equipe" / "Custos"
            nested.mkdir(parents=True)
            expected = nested / "arquivo-automatico.xlsx"
            expected.write_bytes(b"workbook")
            old_roots = os.environ.get("EXCEL_SEARCH_ROOTS")
            old_onedrive = os.environ.get("OneDriveCommercial")
            old_fallback = os.environ.get("EXCEL_FALLBACK_SAMPLE")
            try:
                os.environ.pop("EXCEL_SEARCH_ROOTS", None)
                os.environ["OneDriveCommercial"] = str(onedrive)
                os.environ["EXCEL_FALLBACK_SAMPLE"] = "false"
                located = discover_excel_file("arquivo-automatico.xlsx", base_dir=temp_root)
                self.assertEqual(located, expected.resolve())
            finally:
                if old_roots is None:
                    os.environ.pop("EXCEL_SEARCH_ROOTS", None)
                else:
                    os.environ["EXCEL_SEARCH_ROOTS"] = old_roots
                if old_onedrive is None:
                    os.environ.pop("OneDriveCommercial", None)
                else:
                    os.environ["OneDriveCommercial"] = old_onedrive
                if old_fallback is None:
                    os.environ.pop("EXCEL_FALLBACK_SAMPLE", None)
                else:
                    os.environ["EXCEL_FALLBACK_SAMPLE"] = old_fallback

    def test_template_search_roots_fall_back_to_onedrive_discovery(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            onedrive = temp_root / "OneDrive"
            expected = onedrive / "B2BCUSTOS" / "arquivo-legado.xlsx"
            expected.parent.mkdir(parents=True)
            expected.write_bytes(b"workbook")

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": r"data;C:\Users\SEU_USUARIO\OneDrive - SUA EMPRESA\Pasta Compartilhada",
                    "OneDrive": str(onedrive),
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                located = discover_excel_file("arquivo-legado.xlsx", base_dir=temp_root)

            self.assertEqual(located, expected.resolve())

    def test_legacy_excel_file_name_remains_compatible(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            synced = temp_root / "OneDrive" / "Dados"
            synced.mkdir(parents=True)
            expected = synced / "arquivo-antigo.xlsx"
            expected.write_bytes(b"workbook")

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": str(synced),
                    "EXCEL_FILE_NAME": expected.name,
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_FILENAME", None)
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)

            self.assertEqual(repository.file_path, expected.resolve())

    def test_operational_header_aliases_are_loaded(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            synced = temp_root / "OneDrive" / "B2BCUSTOS"
            synced.mkdir(parents=True)
            workbook_path = synced / "base-operacional.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            sheet.append(
                [
                    "Região",
                    "ID",
                    "Data",
                    "Obra Executada",
                    "Custo MO",
                    "Custo Material",
                    "Custo Total",
                    "Custo Evitado",
                    "Tempo Execução",
                    "Detalhes Obra",
                    "Status Obra",
                    "Vol. Técnicos",
                ]
            )
            sheet.append(
                ["SUL", 665232, "2026-09-05", "REPARO B2B", 810.32, 25, 835.32, 1000, 1, "CLIENTE FINAL", "CONCLUIDO", 2]
            )
            workbook.save(workbook_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": str(synced),
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                items = LocalExcelRepository(temp_root).list_activities()

            self.assertEqual(len(items), 1)
            self.assertEqual(items[0]["tipo_atividade"], "REPARO B2B")
            self.assertEqual(items[0]["status"], "CONCLUIDO")
            self.assertEqual(items[0]["situacao"], "CLIENTE FINAL")
            self.assertEqual(items[0]["custo_total"], 835.32)
            self.assertEqual(items[0]["calculation_mode"], "direct_costs")

    def test_include_preserves_row_style_and_extends_excel_table(self):
        import openpyxl
        from openpyxl.styles import PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            synced = temp_root / "OneDrive"
            synced.mkdir()
            workbook_path = synced / "inclusao.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            headers = [
                "ID", "Data", "Empresa", "Obra Executada", "Custo MO", "Custo Material",
                "Custo Total", "Custo Evitado", "Detalhes Obra", "Status Obra",
                "Matrícula", "Nome do tecnico", "EPS",
            ]
            sheet.append(headers)
            sheet.append([1, "2026-09-01", "VIVO", "REPARO", 100, 50, 150, 250, "CLIENTE", "CONCLUIDO", "A1", "ANA", "VIVO"])
            sheet["A2"].fill = PatternFill("solid", fgColor="EAD1E7")
            table = Table(displayName="AtividadesTable", ref="A1:M2")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            sheet.add_table(table)
            workbook.save(workbook_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": str(synced),
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                created = repository.create_activity(
                    {
                        "data": "2026-09-05",
                        "empresa": "VIVO",
                        "tipo_atividade": "INSTALAÇÃO",
                        "status": "CONCLUIDO",
                        "situacao": "FINALIZADA",
                        "custo_mo": 200,
                        "custo_mat": 75,
                        "custo_evitado": 400,
                        "matricula": "A2",
                        "tecnico_nome": "BRUNO",
                        "eps": "VIVO",
                    }
                )

            verified = openpyxl.load_workbook(workbook_path)
            sheet = verified["Atividades"]
            self.assertEqual(created["id"], "2")
            self.assertEqual(sheet.max_row, 3)
            self.assertEqual(sheet["A3"].fill.fgColor.rgb, sheet["A2"].fill.fgColor.rgb)
            self.assertEqual(sheet.tables["AtividadesTable"].ref, "A1:N3")
            self.assertEqual(sheet["N1"].value, "DRAFT")
            self.assertEqual(sheet["C3"].value, "VIVO")
            self.assertEqual(sheet["G3"].value, 275)
            verified.close()

    def test_settings_catalogs_are_saved_and_drive_form_options(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(workbook_path)

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                historical_statuses = repository.get_options()["statuses"]
                self.assertTrue(historical_statuses)

                saved = repository.save_settings(
                    {
                        "reference_month": "2026-09",
                        "monthly_technician_cost": 15_000,
                        "business_days": 22,
                        "statuses": ["NOVO", "CONCLUÍDO"],
                        "types": ["INSTALAÇÃO", "REPARO"],
                        "technologies": ["GPON", "B2B"],
                        "technicians": [
                            {"registration": "T-001", "name": "Ana Souza"},
                            {"registration": "T-002", "name": "Bruno Lima"},
                        ],
                        "companies": ["Empresa A", "Empresa B"],
                        "eps": ["EPS Norte"],
                    }
                )
                options = repository.get_options()

            self.assertEqual(saved["global_daily_rate"], 681.82)
            self.assertEqual(saved["monthly_technician_cost"], 15_000)
            self.assertEqual(saved["business_days"], 22)
            self.assertEqual(saved["statuses"], ["NOVO", "CONCLUÍDO"])
            self.assertEqual(saved["technicians"][0], {"registration": "T-001", "name": "Ana Souza"})
            self.assertEqual(options["statuses"], ["NOVO", "CONCLUÍDO"])
            self.assertEqual(options["types"], ["INSTALAÇÃO", "REPARO"])
            self.assertEqual(options["technologies"], ["GPON", "B2B"])
            self.assertEqual(options["companies"], ["Empresa A", "Empresa B"])
            self.assertEqual(options["eps"], ["EPS Norte"])
            self.assertNotEqual(options["statuses"], historical_statuses)

            workbook = openpyxl.load_workbook(workbook_path, data_only=True)
            config = workbook["Config"]
            self.assertEqual(config["A8"].value, "Categoria")
            self.assertEqual(config["B8"].value, "Código/Matrícula")
            self.assertEqual(config["C8"].value, "Nome/Valor")
            self.assertEqual(config["B2"].value, 15_000)
            self.assertEqual(config["B3"].value, 22)
            self.assertEqual(config["B4"].value, 681.82)
            categories = [config.cell(row, 1).value for row in range(9, config.max_row + 1)]
            self.assertEqual(categories.count("STATUS"), 2)
            self.assertEqual(categories.count("TIPO_ATIVIDADE"), 2)
            self.assertEqual(categories.count("TECNOLOGIA"), 2)
            self.assertEqual(categories.count("TECNICO"), 2)
            self.assertEqual(categories.count("EMPRESA"), 2)
            self.assertEqual(categories.count("EPS"), 1)
            workbook.close()

    def test_empty_saved_catalog_remains_authoritative(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(workbook_path)

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                repository.save_settings(
                    {
                        "global_daily_rate": 0,
                        "reference_month": "2026-09",
                        "business_days": 22,
                        "statuses": [],
                        "technicians": [],
                        "companies": [],
                        "eps": [],
                    }
                )
                options = repository.get_options()

            self.assertEqual(options["statuses"], [])
            self.assertEqual(options["technicians"], [])
            self.assertEqual(options["companies"], [])
            self.assertEqual(options["eps"], [])
            self.assertTrue(options["types"])

    def test_hidden_modal_fields_are_preserved_when_editing(self):
        import openpyxl

        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "edicao.xlsx"
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Atividades"
            sheet.append(
                [
                    "ID", "Tipo de Atividade", "Status", "Situação", "Material Utilizado",
                    "Cod. Material", "Quantidade", "Custo Mat", "Qtde Serviço", "Custo M.O",
                    "Custo Técnico / Dia", "Qtde Técnicos", "Qtde Dias", "Custo Total",
                    "Custo por Técnico", "Custo Evitado", "GAP",
                ]
            )
            sheet.append([1, "REPARO", "ABERTO", "ORIGINAL", "CABO", "MAT-9", 2, 40, 3, 75, 500, 2, 1, 1155, 577.5, 1300, 145])
            workbook.save(workbook_path)
            workbook.close()

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                repository.update_activity(
                    "1",
                    {
                        "tipo_atividade": "REPARO",
                        "status": "CONCLUÍDO",
                        "situacao": "ATUALIZADA",
                        "quantidade": 2,
                        "custo_mat": 40,
                        "custo_mo": 75,
                        "custo_tecnico_dia": 500,
                        "qtde_tecnicos": 2,
                        "qtde_dias": 1,
                        "custo_evitado": 1300,
                    },
                )

            workbook = openpyxl.load_workbook(workbook_path, data_only=False)
            sheet = workbook["Atividades"]
            self.assertEqual(sheet["E2"].value, "CABO")
            self.assertEqual(sheet["F2"].value, "MAT-9")
            self.assertEqual(sheet["I2"].value, 3)
            self.assertEqual(sheet["N2"].value, "=H2+J2")
            self.assertEqual(sheet["P2"].value, 75)
            self.assertEqual(sheet["Q2"].value, 0)
            workbook.close()

    def test_technology_value_populates_generic_activity_without_daily_team_formula(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(workbook_path)

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                repository.save_settings(
                    {
                        "reference_month": "2026-09",
                        "monthly_technician_cost": 15_000,
                        "business_days": 22,
                        "statuses": ["NOVO"],
                        "types": ["IMPLANTAÇÃO", "REPARO", "ATIVAÇÃO"],
                        "technology_rates": [{"name": "GPON", "service_cost": 1_250}],
                        "category_teams": [
                            {"category": "implantacao", "technicians": 2},
                            {"category": "reparo", "technicians": 1},
                            {"category": "ativacao", "technicians": 3},
                        ],
                        "technicians": [],
                        "companies": [],
                        "eps": [],
                    }
                )
                created = repository.create_activity(
                    {
                        "tipo_atividade": "IMPLANTAÇÃO",
                        "status": "NOVO",
                        "situacao": "TESTE",
                        "tecnologia": "GPON",
                        "custo_mo": 999,
                        "custo_mat": 50,
                        "qtde_tecnicos": 4,
                        "qtde_dias": 3,
                    }
                )

            self.assertEqual(created["custo_mo"], 1_250)
            self.assertEqual(created["custo_total"], 1_300)
            self.assertEqual(created["gap"], 0)

    def test_erb_preserves_service_cost_from_external_calculator(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(workbook_path)

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                repository.save_settings(
                    {
                        "reference_month": "2026-09",
                        "monthly_technician_cost": 15_000,
                        "business_days": 22,
                        "statuses": ["NOVO"],
                        "types": ["IMPLANTAÇÃO"],
                        "technology_rates": [{"name": "ERB", "service_cost": 1_250}],
                        "category_teams": [],
                        "technicians": [],
                        "companies": [],
                        "eps": [],
                    }
                )
                created = repository.create_activity(
                    {
                        "tipo_atividade": "IMPLANTAÇÃO",
                        "status": "NOVO",
                        "situacao": "CALCULADORA EXTERNA",
                        "tecnologia": "ERB",
                        "custo_mo": "777,50",
                        "custo_mat": "22,50",
                    }
                )

            self.assertEqual(created["custo_mo"], 777.5)
            self.assertEqual(created["custo_evitado"], 777.5)
            self.assertEqual(created["custo_total"], 800)

    def test_reference_workbook_upload_is_saved_beside_main_excel(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            temp_root = Path(temp_dir)
            data_dir = temp_root / "data"
            data_dir.mkdir()
            workbook_path = data_dir / "B2B_CTACUSTOS.xlsx"
            create_test_activity_workbook(workbook_path)
            source = workbook_path

            with patch.dict(
                os.environ,
                {
                    "EXCEL_SEARCH_ROOTS": "data",
                    "EXCEL_FILENAME": workbook_path.name,
                    "EXCEL_SHEET_NAME": "Atividades",
                    "EXCEL_FALLBACK_SAMPLE": "false",
                    "REFERENCE_UPLOAD_DIR": "",
                },
                clear=False,
            ):
                os.environ.pop("EXCEL_PATH", None)
                repository = LocalExcelRepository(temp_root)
                uploaded = repository.save_reference_workbook("services", "servicos.xlsx", source.read_bytes())
                settings = repository.get_settings()

            uploaded_path = Path(uploaded["path"])
            self.assertTrue(uploaded_path.is_file())
            self.assertEqual(uploaded_path.parent, data_dir / "B2B_CTACUSTOS_Importacoes")
            self.assertEqual(settings["uploads"]["services"]["filename"], uploaded_path.name)
            self.assertEqual(settings["uploads"]["services"]["path"], str(uploaded_path))


if __name__ == "__main__":
    unittest.main()
