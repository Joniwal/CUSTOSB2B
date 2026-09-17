import io
import json
import os
import sys
import tempfile
import threading
import unittest
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch
from urllib.request import urlopen

PROJECT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT))

from app import B2BServer, RequestHandler  # noqa: E402
from excel_repository import _onedrive_registry_roots  # noqa: E402


class ApiTests(unittest.TestCase):
    def setUp(self):
        import openpyxl
        from openpyxl.styles import PatternFill
        from openpyxl.worksheet.table import Table, TableStyleInfo

        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.path = Path(self.temp.name) / "B2B_CTACUSTOS.xlsx"
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Atividades"
        sheet.append([
            "ID", "Tipo de Atividade", "Status", "Situação", "Tecnologia", "Data",
            "Custo MO", "Custo Material", "Custo Total", "Custo Técnico / Dia",
            "Qtde Técnicos", "Qtde Dias", "DRAFT",
        ])
        for row in (
            [1, "Reparo", "Concluída", "Teste", "GPON", "2026-09-03", 300, 30, 330, 100, 1, 1, "RASCUNHO-1"],
            [2, "Ativação", "Concluída", "Teste", "B2B", "2026-09-04", 400, 40, 440, 100, 1, 1, "RASCUNHO-2"],
            [3, "Reparo", "Concluída", "Teste", "GPON", "2026-08-03", 150, 15, 165, 100, 1, 1, ""],
            [4, "Ativação", "Concluída", "Teste", "B2B", "2026-08-04", 200, 20, 220, 100, 1, 1, ""],
        ):
            sheet.append(row)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="DD7FD4")
        sheet["G2"].number_format = '"R$" #,##0.00'
        table = Table(displayName="Atividades", ref="A1:M5")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium4", showRowStripes=True)
        sheet.add_table(table)
        workbook.save(self.path)
        workbook.close()
        self.env_patch = patch.dict(os.environ, {
            "EXCEL_PATH": str(self.path), "EXCEL_SHEET_NAME": "Atividades",
            "EXCEL_FALLBACK_SAMPLE": "false",
        })
        self.env_patch.start()
        self.addCleanup(self.env_patch.stop)
        self.server = B2BServer(("127.0.0.1", 0), RequestHandler)
        self.thread = threading.Thread(target=self.server.serve_forever, daemon=True)
        self.thread.start()
        self.addCleanup(self.stop_server)
        self.url = f"http://127.0.0.1:{self.server.server_port}"

    def stop_server(self):
        self.server.shutdown()
        self.thread.join(timeout=3)
        self.server.server_close()

    def get_json(self, route):
        with urlopen(self.url + route, timeout=5) as response:
            return json.load(response)

    def test_dashboard_technology_filters_current_previous_and_summaries(self):
        data = self.get_json("/api/dashboard?month=2026-09&technology=gpon")
        self.assertEqual(data["metrics"]["total"], 1)
        labor = next(item for item in data["kpis"] if item["key"] == "custo_mo")
        self.assertEqual((labor["current"], labor["previous"]), (300, 150))
        self.assertEqual(data["period"]["previous_start"], "2026-08-01")
        self.assertEqual(data["charts"]["by_type"][0]["labor"], 300)
        self.assertEqual(data["charts"]["by_type"][0]["gap"], 200)
        self.assertEqual([item["label"] for item in data["charts"]["by_technology"]], ["GPON"])
        self.assertEqual(data["filters"]["technologies"], ["B2B", "GPON"])
        categories = {item["key"]: item for item in data["category_kpis"]}
        self.assertEqual(categories["reparo"]["service"], 300)
        self.assertEqual(categories["reparo"]["material"], 30)
        self.assertEqual(categories["reparo"]["gap"], 300)
        self.assertEqual(categories["reparo"]["trends"]["service"]["previous"], 150)

    def test_list_technology_and_draft_search_and_sort(self):
        data = self.get_json("/api/activities?technology=GPON&q=RASCUNHO-1&sort=draft")
        self.assertEqual(data["pagination"]["total"], 1)
        self.assertEqual(data["items"][0]["draft"], "RASCUNHO-1")

    def test_export_download_omits_draft_and_does_not_modify_source(self):
        import openpyxl

        before = self.path.read_bytes()
        with urlopen(self.url + "/api/activities/export", timeout=5) as response:
            self.assertIn("attachment", response.headers["Content-Disposition"])
            self.assertEqual(response.headers["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            exported = openpyxl.load_workbook(io.BytesIO(response.read()))
        sheet = exported["Atividades"]
        self.assertEqual(sheet.max_column, 12)
        self.assertEqual(sheet.max_row, 5)
        self.assertNotIn("DRAFT", [cell.value for cell in sheet[1]])
        self.assertEqual(sheet["G2"].number_format, '"R$" #,##0.00')
        self.assertEqual(sheet["A1"].fill.fgColor.rgb, "00DD7FD4")
        self.assertEqual(sheet.tables["Atividades"].ref, "A1:L5")
        exported.close()
        self.assertEqual(self.path.read_bytes(), before)


class WindowsDiscoveryTests(unittest.TestCase):
    @unittest.skipUnless(os.name == "nt", "Windows registry discovery")
    def test_registered_sharepoint_library_outside_onedrive_is_discovered(self):
        accounts = r"Software\Microsoft\OneDrive\Accounts"
        business = accounts + r"\Business1"
        cache = business + r"\ScopeIdToMountPointPathCache"
        expected = Path(r"C:\Users\TESTE\Empresa\Equipe - Documentos")
        children = {accounts: ["Business1"], business: ["ScopeIdToMountPointPathCache"], cache: []}

        def open_key(_hive, key_path):
            if key_path not in children:
                raise OSError("Key not present")
            return nullcontext(key_path)

        def query_value(key_path, name):
            raise OSError("Value not present")

        def enum_key(key_path, index):
            try:
                return children[key_path][index]
            except IndexError:
                raise OSError("End of keys")

        def enum_value(key_path, index):
            if key_path == cache and index == 0:
                return "scope-id", str(expected), 1
            raise OSError("End of values")

        fake_registry = SimpleNamespace(
            HKEY_CURRENT_USER=1, OpenKey=open_key, QueryValueEx=query_value,
            EnumKey=enum_key, EnumValue=enum_value,
        )
        with patch.dict(sys.modules, {"winreg": fake_registry}):
            self.assertIn(expected, _onedrive_registry_roots())


if __name__ == "__main__":
    unittest.main()
